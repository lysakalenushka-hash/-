#!/usr/bin/env python3
"""Раздел I Приложения 2 — вопросы электросетевых организаций (актуализация 2026)."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC_APP = Path("/home/ubuntu/.cursor/projects/workspace/uploads/___________2_sent_891c.xlsx")
SRC_CMP = Path("/home/ubuntu/.cursor/projects/workspace/uploads/Voprosi-po-elektrobezopasnosti-2026_d018.xlsx")
OUT_DIR = Path("/workspace/elektrobezopasnost_razdel_I")
OUT = OUT_DIR / "Elektrobezopasnost_razdel_I.xlsx"
OUT_CYR = OUT_DIR / "Электробезопасность_раздел_I.xlsx"

HEAD = PatternFill("solid", fgColor="1F4E79")
TOPIC = PatternFill("solid", fgColor="2E75B6")
OK = PatternFill("solid", fgColor="C6EFCE")
ALT = PatternFill("solid", fgColor="D6EAF8")
WARN = PatternFill("solid", fgColor="FFF2CC")
DEL = PatternFill("solid", fgColor="F8CBAD")
ADD = PatternFill("solid", fgColor="C6EFCE")
CHG = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HF = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TF = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
CF = Font(name="Calibri", size=11)
TITLE = Font(name="Calibri", bold=True, size=16, color="1F4E79")


def parse_section_i(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb["I"]
    topics = []
    topic = None
    questions = []
    r, maxr = 3, ws.max_row
    while r <= maxr:
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        flags = [ws.cell(r, c).value for c in range(3, 9)]
        is_topic = a and not b and all(x is None for x in flags)
        if is_topic:
            topic = str(a).strip()
            topics.append(topic)
            r += 1
            continue
        num = None
        if a is not None and b:
            try:
                num = int(str(a).strip())
            except ValueError:
                num = str(a).strip()
            q = {
                "topic": topic,
                "n": num,
                "text": str(b).strip(),
                "flags": flags,
                "opts": [],
            }
            r += 1
            while r <= maxr:
                a2, b2 = ws.cell(r, 1).value, ws.cell(r, 2).value
                f2 = [ws.cell(r, c).value for c in range(3, 9)]
                if a2 is not None and b2:
                    break
                if a2 and not b2 and all(x is None for x in f2):
                    break
                if b2:
                    q["opts"].append(str(b2).strip())
                    r += 1
                    continue
                r += 1
                break
            questions.append(q)
            continue
        r += 1
    wb.close()
    return topics, questions


def flag_text(flags):
    names = ["до 1000 В", "до и выше 1000 В", "II", "III", "IV", "V"]
    on = [n for n, v in zip(names, flags) if v]
    volt = [x for x in on if "В" in x]
    grp = [x for x in on if x in ("II", "III", "IV", "V")]
    return ", ".join(volt), ", ".join(grp)


def style_head(cell, text):
    cell.value = text
    cell.fill = HEAD
    cell.font = HF
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN


def style_cell(cell, val, wrap=True):
    cell.value = val
    cell.font = CF
    cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="left")
    cell.border = THIN


def build():
    topics, questions = parse_section_i(SRC_APP)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ---- сводка ----
    w0 = wb.active
    w0.title = "0. Сводка"
    w0["A1"] = "Приложение 2. Раздел I — работники электросетевых (обслуживающих) организаций"
    w0["A1"].font = TITLE
    w0.merge_cells("A1:F1")
    w0.row_dimensions[1].height = 28
    w0["A2"] = (
        "АКТУАЛЬНО с 01.09.2026. Перечень утверждён Ростехнадзором 10.08.2026, "
        "в ЕПТ применяется с 1 сентября 2026 г. Раздел I сверен с официальным текстом: "
        "569 из 569 вопросов совпадают (разница только «г.» в дате приказа Минэнерго № 1070). "
        "Правильный ответ в компоновке Приложения 2 — первый вариант. "
        "Знаки «+» — область применения (напряжение и группа II–V), не ключ."
    )
    w0["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    w0.merge_cells("A2:F2")
    w0.row_dimensions[2].height = 48

    for i, h in enumerate(["№", "Нормативный документ (тема)", "Вопросов", "Сквозные №", "Напряжение (типично)", "Группы (типично)"], 1):
        style_head(w0.cell(4, i), h)
    w0.row_dimensions[4].height = 24

    from collections import defaultdict
    by = defaultdict(list)
    for i, q in enumerate(questions, 1):
        by[q["topic"]].append((i, q))
    r = 5
    start = 1
    for t in topics:
        qs = by[t]
        end = start + len(qs) - 1
        vset, gset = set(), set()
        for _, q in qs:
            v, g = flag_text(q["flags"])
            if v:
                vset.add(v)
            if g:
                gset.add(g)
        vals = [r - 4, t, len(qs), f"{start}–{end}", "; ".join(sorted(vset)), "; ".join(sorted(gset))]
        for c, v in enumerate(vals, 1):
            style_cell(w0.cell(r, c), v)
            if r % 2 == 0:
                w0.cell(r, c).fill = ALT
        w0.row_dimensions[r].height = 32
        start = end + 1
        r += 1
    style_cell(w0.cell(r, 1), "ИТОГО")
    w0.cell(r, 1).font = Font(name="Calibri", bold=True, size=11)
    style_cell(w0.cell(r, 2), "")
    style_cell(w0.cell(r, 3), len(questions))
    w0.cell(r, 3).font = Font(name="Calibri", bold=True, size=11)
    for c in range(4, 7):
        style_cell(w0.cell(r, c), "")
    for c in range(1, 7):
        w0.cell(r, c).fill = WARN
    w0.column_dimensions["A"].width = 6
    w0.column_dimensions["B"].width = 78
    w0.column_dimensions["C"].width = 12
    w0.column_dimensions["D"].width = 14
    w0.column_dimensions["E"].width = 28
    w0.column_dimensions["F"].width = 18
    w0.freeze_panes = "A5"
    w0.page_setup.orientation = "landscape"
    w0.page_setup.fitToPage = True
    w0.page_setup.fitToWidth = 1
    w0.page_setup.fitToHeight = 1

    # ---- вопросы ----
    w1 = wb.create_sheet("1. Вопросы")
    w1["A1"] = "Раздел I. Вопросы и варианты ответов (верный — вариант 1)"
    w1["A1"].font = TITLE
    w1.merge_cells("A1:L1")
    w1.row_dimensions[1].height = 24
    headers = [
        "Скв. №",
        "Тема",
        "№ в теме",
        "Вопрос",
        "Верный ответ (вар. 1)",
        "Вариант 2",
        "Вариант 3",
        "Вариант 4",
        "Вариант 5+",
        "Напряжение",
        "Группа II–V",
        "Число вариантов",
    ]
    for i, h in enumerate(headers, 1):
        style_head(w1.cell(2, i), h)
    w1.row_dimensions[2].height = 32
    w1.auto_filter.ref = f"A2:L{2 + len(questions)}"
    w1.freeze_panes = "A3"

    for i, q in enumerate(questions, 1):
        rr = 2 + i
        v, g = flag_text(q["flags"])
        extra = " | ".join(q["opts"][4:]) if len(q["opts"]) > 4 else ""
        opts = q["opts"] + ["", "", "", ""]
        vals = [
            i,
            q["topic"],
            q["n"],
            q["text"],
            opts[0],
            opts[1],
            opts[2],
            opts[3],
            extra,
            v,
            g,
            len(q["opts"]),
        ]
        for c, val in enumerate(vals, 1):
            style_cell(w1.cell(rr, c), val)
        w1.cell(rr, 5).fill = OK
        w1.cell(rr, 5).font = Font(name="Calibri", size=11, bold=True)
        if i % 2 == 0:
            for c in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12):
                if w1.cell(rr, c).fill.fgColor is None or c != 5:
                    w1.cell(rr, c).fill = ALT
            w1.cell(rr, 5).fill = OK
        w1.row_dimensions[rr].height = 48

    widths = [10, 42, 10, 55, 42, 38, 38, 38, 28, 22, 14, 12]
    for i, w in enumerate(widths, 1):
        w1.column_dimensions[get_column_letter(i)].width = w
    w1.page_setup.orientation = "landscape"
    w1.page_setup.fitToPage = True
    w1.page_setup.fitToWidth = 1
    w1.page_setup.fitToHeight = 0
    w1.print_title_rows = "1:2"

    # ---- тренажёр: блоки ----
    w2 = wb.create_sheet("2. Тренажёр")
    w2["A1"] = "Раздел I. Карточки: вопрос + варианты (зелёный — верный)"
    w2["A1"].font = TITLE
    w2.merge_cells("A1:C1")
    row = 3
    letters = "АБВГДЕЖЗИК"
    for i, q in enumerate(questions, 1):
        w2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = w2.cell(row, 1, q["topic"])
        cell.fill = TOPIC
        cell.font = TF
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        w2.row_dimensions[row].height = 20
        row += 1
        w2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        v, g = flag_text(q["flags"])
        head = f"{i}.  [{q['n']}]  {q['text']}"
        cell = w2.cell(row, 1, head)
        cell.font = Font(name="Calibri", bold=True, size=12)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        w2.row_dimensions[row].height = 36
        row += 1
        w2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        meta = w2.cell(row, 1, f"Напряжение: {v or '—'}    ·    Группа: {g or '—'}")
        meta.font = Font(name="Calibri", italic=True, size=10, color="666666")
        row += 1
        for j, opt in enumerate(q["opts"]):
            w2.cell(row, 1, letters[j] if j < len(letters) else str(j + 1)).border = THIN
            w2.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
            w2.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            oc = w2.cell(row, 2, opt)
            oc.font = CF
            oc.alignment = Alignment(wrap_text=True, vertical="center")
            oc.border = THIN
            w2.cell(row, 3).border = THIN
            if j == 0:
                w2.cell(row, 1).fill = OK
                oc.fill = OK
                oc.font = Font(name="Calibri", size=11, bold=True)
            w2.row_dimensions[row].height = max(22, 14 + 12 * (len(opt) // 90))
            row += 1
        row += 1
        if i % 50 == 0:
            print("trainer", i)
    w2.column_dimensions["A"].width = 8
    w2.column_dimensions["B"].width = 90
    w2.column_dimensions["C"].width = 20
    w2.page_setup.orientation = "portrait"
    w2.page_setup.fitToPage = True
    w2.page_setup.fitToWidth = 1
    w2.page_setup.fitToHeight = 0
    w2.freeze_panes = "A3"

    # ---- изменения 2026, только раздел I ----
    w3 = wb.create_sheet("3. Что изменилось")
    cmp = load_workbook(SRC_CMP, data_only=True)
    src = cmp["Сравнение"]
    w3["A1"] = "Раздел I. Сопоставление со старым банком (файл Voprosi-po-elektrobezopasnosti-2026)"
    w3["A1"].font = TITLE
    w3.merge_cells("A1:K1")
    w3.row_dimensions[1].height = 24
    for c in range(1, 12):
        style_head(w3.cell(2, c), src.cell(1, c).value)
    w3.row_dimensions[2].height = 30
    rr = 3
    counts = {}
    for r in range(2, src.max_row + 1):
        if src.cell(r, 1).value != "I":
            continue
        status = src.cell(r, 2).value
        counts[status] = counts.get(status, 0) + 1
        for c in range(1, 12):
            style_cell(w3.cell(rr, c), src.cell(r, c).value)
        fill = {"Удалён": DEL, "Добавлен": ADD, "Изменена формулировка": CHG, "Без изменений": None}.get(status)
        if fill:
            for c in range(1, 12):
                w3.cell(rr, c).fill = fill
        w3.row_dimensions[rr].height = 36
        rr += 1
    cmp.close()
    w3.auto_filter.ref = f"A2:K{rr - 1}"
    w3.freeze_panes = "A3"
    w3.column_dimensions["A"].width = 10
    w3.column_dimensions["B"].width = 24
    w3.column_dimensions["C"].width = 12
    w3.column_dimensions["D"].width = 45
    w3.column_dimensions["E"].width = 12
    w3.column_dimensions["F"].width = 45
    w3.column_dimensions["G"].width = 36
    w3.column_dimensions["H"].width = 36
    w3.column_dimensions["I"].width = 28
    w3.column_dimensions["J"].width = 22
    w3.column_dimensions["K"].width = 12
    w3.page_setup.orientation = "landscape"
    w3.page_setup.fitToPage = True
    w3.page_setup.fitToWidth = 1
    w3.page_setup.fitToHeight = 0
    w3.print_title_rows = "1:2"

    # counts on summary
    w0["A20"] = "Изменения относительно предыдущего банка (раздел I)"
    w0["A20"].font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    w0.merge_cells("A20:C20")
    for i, h in enumerate(["Статус", "Количество", ""], 1):
        if h:
            style_head(w0.cell(21, i), h)
    r = 22
    for k in ("Без изменений", "Изменена формулировка", "Удалён", "Добавлен"):
        style_cell(w0.cell(r, 1), k)
        style_cell(w0.cell(r, 2), counts.get(k, 0))
        if k == "Удалён":
            w0.cell(r, 1).fill = DEL
            w0.cell(r, 2).fill = DEL
        elif k == "Добавлен":
            w0.cell(r, 1).fill = ADD
            w0.cell(r, 2).fill = ADD
        elif k == "Изменена формулировка":
            w0.cell(r, 1).fill = CHG
            w0.cell(r, 2).fill = CHG
        r += 1

    wb.save(OUT)
    shutil.copyfile(OUT, OUT_CYR)
    print("questions", len(questions), "topics", len(topics), "cmp", counts)
    print("saved", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
