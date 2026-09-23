#!/usr/bin/env python3
"""Раздел II Приложения 2 — электростанции / ГЭС (актуализация 2026)."""

from __future__ import annotations

import shutil
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC_APP = Path("/home/ubuntu/.cursor/projects/workspace/uploads/___________2_sent_891c.xlsx")
SRC_CMP = Path("/home/ubuntu/.cursor/projects/workspace/uploads/Voprosi-po-elektrobezopasnosti-2026_d018.xlsx")
OUT_DIR = Path("/workspace/elektrobezopasnost_razdel_II")
OUT = OUT_DIR / "Elektrobezopasnost_razdel_II.xlsx"
OUT_CYR = OUT_DIR / "Электробезопасность_раздел_II.xlsx"
OUT_LMS = OUT_DIR / "Elektrobezopasnost_razdel_II_LMS.xlsx"
OUT_LMS_CYR = OUT_DIR / "Электробезопасность_раздел_II_тест.xlsx"

FOLDER = "Электробезопасность. Раздел II"
PROGRAM = "Электробезопасность. Раздел II"

HEAD = PatternFill("solid", fgColor="1F4E79")
OK = PatternFill("solid", fgColor="C6EFCE")
ALT = PatternFill("solid", fgColor="D6EAF8")
WARN = PatternFill("solid", fgColor="FFF2CC")
DEL = PatternFill("solid", fgColor="F8CBAD")
ADD = PatternFill("solid", fgColor="C6EFCE")
CHG = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HF = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CF = Font(name="Calibri", size=11)
TITLE = Font(name="Calibri", bold=True, size=16, color="1F4E79")


def format_option(cell) -> str:
    val = cell.value
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val)
    if isinstance(val, (int, float)):
        fmt = (cell.number_format or "").lower()
        if "%" in fmt:
            pct = round(float(val) * 100, 10)
            if pct == int(pct):
                return f"{int(pct)}%"
            return f"{pct:g}%"
        if float(val) == int(val):
            return str(int(val))
        return str(val).strip()
    return str(val).strip()


def parse_section_ii(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb["II"]
    topics = []
    topic = None
    questions = []
    r, maxr = 3, ws.max_row
    while r <= maxr:
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        flags = [ws.cell(r, c).value for c in range(3, 9)]
        is_topic = a and not b and all(x is None for x in flags)
        if is_topic:
            topic = str(a).strip()
            topics.append(topic)
            r += 1
            continue
        if a is not None and b:
            try:
                num = int(str(a).strip())
            except ValueError:
                num = str(a).strip()
            q = {
                "topic": topic,
                "n": num,
                "text": str(b).strip(),
                "flags": flags,
                "opts": [],
            }
            r += 1
            while r <= maxr:
                a2 = ws.cell(r, 1).value
                cell_b = ws.cell(r, 2)
                b2 = cell_b.value
                f2 = [ws.cell(r, c).value for c in range(3, 9)]
                if a2 is not None and b2:
                    break
                if a2 and not b2 and all(x is None for x in f2):
                    break
                if b2 is not None and b2 != "":
                    q["opts"].append(format_option(cell_b))
                    r += 1
                    continue
                r += 1
                break
            questions.append(q)
            continue
        r += 1
    wb.close()
    return topics, questions


def flag_text(flags):
    names = ["до 1000 В", "до и выше 1000 В", "II", "III", "IV", "V"]
    on = [n for n, v in zip(names, flags) if v]
    volt = [x for x in on if "В" in x]
    grp = [x for x in on if x in ("II", "III", "IV", "V")]
    return ", ".join(volt), ", ".join(grp)


def style_head(cell, text):
    cell.value = text
    cell.fill = HEAD
    cell.font = HF
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN


def style_cell(cell, val, wrap=True):
    cell.value = val
    cell.font = CF
    cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="left")
    cell.border = THIN


def build(topics, questions):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    w0 = wb.active
    w0.title = "0. Сводка"
    w0["A1"] = (
        "Приложение 2. Раздел II — работники электростанций, "
        "в том числе комбинированной выработки и гидроэлектростанций"
    )
    w0["A1"].font = TITLE
    w0.merge_cells("A1:F1")
    w0.row_dimensions[1].height = 32
    w0["A2"] = (
        "АКТУАЛЬНО с 01.09.2026. Перечень утверждён Ростехнадзором 10.08.2026 "
        "(врио начальника УГЭН Селехов М.В.), в ЕПТ применяется с 1 сентября 2026 г. "
        "Раздел II сверен с официальным текстом База НПА: 841 из 841 вопросов совпадают. "
        "Расхождения только оформительские (пробел после «согласно», «г.» / «№» / кавычки, °C/°С, мм2/мм²). "
        "В исходном xlsx у вопроса ПТЭЭСС № 299 проценты были числами Excel (0,1); в выгрузке восстановлено 10%/15%/17%/20% по официальному перечню. "
        "Правильный ответ в компоновке Приложения 2 — первый вариант. Знаки «+» — область применения, не ключ."
    )
    w0["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    w0.merge_cells("A2:F2")
    w0.row_dimensions[2].height = 72

    for i, h in enumerate(
        ["№", "Нормативный документ (тема)", "Вопросов", "Сквозные №", "Напряжение (типично)", "Группы (типично)"],
        1,
    ):
        style_head(w0.cell(4, i), h)
    w0.row_dimensions[4].height = 24

    by = defaultdict(list)
    for i, q in enumerate(questions, 1):
        by[q["topic"]].append((i, q))
    r = 5
    start = 1
    for t in topics:
        qs = by[t]
        end = start + len(qs) - 1
        vset, gset = set(), set()
        for _, q in qs:
            v, g = flag_text(q["flags"])
            if v:
                vset.add(v)
            if g:
                gset.add(g)
        vals = [r - 4, t, len(qs), f"{start}–{end}", "; ".join(sorted(vset)), "; ".join(sorted(gset))]
        for c, v in enumerate(vals, 1):
            style_cell(w0.cell(r, c), v)
            if r % 2 == 0:
                w0.cell(r, c).fill = ALT
        w0.row_dimensions[r].height = 32
        start = end + 1
        r += 1
    style_cell(w0.cell(r, 1), "ИТОГО")
    w0.cell(r, 1).font = Font(name="Calibri", bold=True, size=11)
    style_cell(w0.cell(r, 2), "")
    style_cell(w0.cell(r, 3), len(questions))
    w0.cell(r, 3).font = Font(name="Calibri", bold=True, size=11)
    for c in range(4, 7):
        style_cell(w0.cell(r, c), "")
    for c in range(1, 7):
        w0.cell(r, c).fill = WARN
    w0.column_dimensions["A"].width = 6
    w0.column_dimensions["B"].width = 78
    w0.column_dimensions["C"].width = 12
    w0.column_dimensions["D"].width = 14
    w0.column_dimensions["E"].width = 28
    w0.column_dimensions["F"].width = 18
    w0.freeze_panes = "A5"

    w1 = wb.create_sheet("1. Вопросы")
    w1["A1"] = "Раздел II. Вопросы и варианты ответов (верный — вариант 1)"
    w1["A1"].font = TITLE
    w1.merge_cells("A1:L1")
    headers = [
        "Скв. №",
        "Тема",
        "№ в теме",
        "Вопрос",
        "Верный ответ (вар. 1)",
        "Вариант 2",
        "Вариант 3",
        "Вариант 4",
        "Вариант 5+",
        "Напряжение",
        "Группа II–V",
        "Число вариантов",
    ]
    for i, h in enumerate(headers, 1):
        style_head(w1.cell(2, i), h)
    w1.row_dimensions[2].height = 32
    w1.auto_filter.ref = f"A2:L{2 + len(questions)}"
    w1.freeze_panes = "A3"

    for i, q in enumerate(questions, 1):
        rr = 2 + i
        v, g = flag_text(q["flags"])
        extra = " | ".join(q["opts"][4:]) if len(q["opts"]) > 4 else ""
        opts = q["opts"] + ["", "", "", ""]
        vals = [i, q["topic"], q["n"], q["text"], opts[0], opts[1], opts[2], opts[3], extra, v, g, len(q["opts"])]
        for c, val in enumerate(vals, 1):
            style_cell(w1.cell(rr, c), val)
        w1.cell(rr, 5).fill = OK
        w1.cell(rr, 5).font = Font(name="Calibri", size=11, bold=True)
        if i % 2 == 0:
            for c in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12):
                w1.cell(rr, c).fill = ALT
            w1.cell(rr, 5).fill = OK
        w1.row_dimensions[rr].height = 48

    widths = [10, 42, 10, 55, 42, 38, 38, 38, 28, 22, 14, 12]
    for i, w in enumerate(widths, 1):
        w1.column_dimensions[get_column_letter(i)].width = w

    w3 = wb.create_sheet("2. Что изменилось")
    cmp = load_workbook(SRC_CMP, data_only=True)
    src = cmp["Сравнение"]
    w3["A1"] = "Раздел II. Сопоставление со старым банком (файл Voprosi-po-elektrobezopasnosti-2026)"
    w3["A1"].font = TITLE
    w3.merge_cells("A1:K1")
    for c in range(1, 12):
        style_head(w3.cell(2, c), src.cell(1, c).value)
    rr = 3
    counts = {}
    for r in range(2, src.max_row + 1):
        if src.cell(r, 1).value != "II":
            continue
        status = src.cell(r, 2).value
        counts[status] = counts.get(status, 0) + 1
        for c in range(1, 12):
            style_cell(w3.cell(rr, c), src.cell(r, c).value)
        fill = {"Удалён": DEL, "Добавлен": ADD, "Изменена формулировка": CHG, "Без изменений": None}.get(status)
        if fill:
            for c in range(1, 12):
                w3.cell(rr, c).fill = fill
        w3.row_dimensions[rr].height = 36
        rr += 1
    cmp.close()
    w3.auto_filter.ref = f"A2:K{rr - 1}"
    w3.freeze_panes = "A3"
    for col, w in zip("ABCDEFGHIJK", [10, 24, 12, 45, 12, 45, 36, 36, 28, 22, 12]):
        w3.column_dimensions[col].width = w

    w0["A20"] = "Изменения относительно предыдущего банка (раздел II)"
    w0["A20"].font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    w0.merge_cells("A20:C20")
    style_head(w0.cell(21, 1), "Статус")
    style_head(w0.cell(21, 2), "Количество")
    r = 22
    for k in ("Без изменений", "Изменена формулировка", "Удалён", "Добавлен"):
        style_cell(w0.cell(r, 1), k)
        style_cell(w0.cell(r, 2), counts.get(k, 0))
        if k == "Удалён":
            w0.cell(r, 1).fill = DEL
            w0.cell(r, 2).fill = DEL
        elif k == "Добавлен":
            w0.cell(r, 1).fill = ADD
            w0.cell(r, 2).fill = ADD
        elif k == "Изменена формулировка":
            w0.cell(r, 1).fill = CHG
            w0.cell(r, 2).fill = CHG
        r += 1

    w4 = wb.create_sheet("3. Актуальность РТН")
    w4["A1"] = "Проверка актуальности раздела II относительно официального перечня РТН"
    w4["A1"].font = TITLE
    w4.merge_cells("A1:B1")
    w4["A2"] = (
        "Источник: перечень утв. 10.08.2026, применяется с 01.09.2026 "
        "(КонсультантПлюс LAW_541956; База НПА h7217878). "
        "Сверка 23.09.2026: 841/841 вопросов раздела II совпадают с официальным текстом. "
        "Содержательных расхождений нет."
    )
    w4["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    w4.merge_cells("A2:B2")
    w4.row_dimensions[2].height = 64
    for i, h in enumerate(["Показатель", "Значение"], 1):
        style_head(w4.cell(4, i), h)
    rows = [
        ("Раздел официального перечня", "II. Работники электростанций, комбинированной выработки и ГЭС"),
        ("Вопросов в Приложении 2 (лист II)", len(questions)),
        ("Вопросов в официальном перечне", 841),
        ("Совпадение формулировок", "841 / 841"),
        ("Тем (НПА)", len(topics)),
        ("Старый банк (до 01.09.2026)", "833 вопроса"),
        ("Без изменений vs старый банк", counts.get("Без изменений", 0)),
        ("Изменена формулировка vs старый банк", counts.get("Изменена формулировка", 0)),
        ("Удалено из старого банка", counts.get("Удалён", 0)),
        ("Добавлено в новый банк", counts.get("Добавлен", 0)),
    ]
    for i, (a, b) in enumerate(rows, 5):
        style_cell(w4.cell(i, 1), a)
        style_cell(w4.cell(i, 2), b)
        if i % 2 == 0:
            w4.cell(i, 1).fill = ALT
            w4.cell(i, 2).fill = ALT
    w4.column_dimensions["A"].width = 48
    w4.column_dimensions["B"].width = 72

    wb.save(OUT)
    shutil.copyfile(OUT, OUT_CYR)
    print("questions", len(questions), "topics", len(topics), "cmp", counts)
    print("saved", OUT, OUT.stat().st_size)
    return counts


def build_lms(questions):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    headers = [
        "Папка",
        "Номер вопроса",
        "Тип записи",
        "Текст вопроса/ответа",
        "Минимальное количество правильных ответов / Правильный или нет",
        "Тип ответа",
        "Комментарий",
        "Изображение",
        "Теги",
    ]
    wrap_head = Alignment(wrap_text=True, vertical="center")
    wrap_body = Alignment(wrap_text=True, vertical="center")
    hf = Font(name="Calibri", bold=True, size=11)
    bf = Font(name="Calibri", size=11)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = hf
        cell.alignment = wrap_head
    ws.row_dimensions[1].height = 48

    r = 2
    for i, q in enumerate(questions, 1):
        ws.cell(r, 1, FOLDER).font = bf
        ws.cell(r, 1).alignment = wrap_body
        ws.cell(r, 2, i).font = bf
        ws.cell(r, 3, "Вопрос").font = bf
        ws.cell(r, 4, q["text"]).font = bf
        ws.cell(r, 4).alignment = wrap_body
        ws.cell(r, 5, 1).font = bf
        ws.cell(r, 6, "Выбор ответа").font = bf
        ws.cell(r, 9, FOLDER).font = bf
        ws.cell(r, 9).alignment = wrap_body
        r += 1
        for j, opt in enumerate(q["opts"]):
            ws.cell(r, 1, FOLDER).font = bf
            ws.cell(r, 1).alignment = wrap_body
            ws.cell(r, 2, i).font = bf
            ws.cell(r, 3, "Ответ").font = bf
            ws.cell(r, 4, opt).font = bf
            ws.cell(r, 4).alignment = wrap_body
            ws.cell(r, 5, 1 if j == 0 else 0).font = bf
            ws.cell(r, 9).alignment = wrap_body
            r += 1

    for col, w in {"A": 28, "B": 14, "C": 14, "D": 72, "E": 24, "F": 18, "G": 28, "H": 16, "I": 26}.items():
        ws.column_dimensions[col].width = w

    wp = wb.create_sheet("Теги программ")
    wp["A1"] = "Наименование программ"
    wp["B1"] = "Уникальные теги"
    wp["C1"] = "Доп. Теги"
    wp["A2"] = PROGRAM
    wp["B2"] = FOLDER
    wo = wb.create_sheet("Теги ОКВЭД")
    wo["A1"] = "Наименование ОКВЭД"
    wo["B1"] = "Уникальные теги"
    wk = wb.create_sheet("Теги контингентов")
    wk["A1"] = "Наименование контингента"
    wk["B1"] = "Уникальные теги"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_LMS)
    shutil.copyfile(OUT_LMS, OUT_LMS_CYR)
    print("lms rows", r - 1, "saved", OUT_LMS, OUT_LMS.stat().st_size)


if __name__ == "__main__":
    topics, questions = parse_section_ii(SRC_APP)
    build(topics, questions)
    build_lms(questions)
