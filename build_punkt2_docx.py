#!/usr/bin/env python3
"""Пункт 2 по методичке: таблица на каждый реактор каждой системы."""

from __future__ import annotations

import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path("/workspace/reaktory_sistemy_variant2")
CA0 = 30.0
K1, K2 = 0.6, 0.4

# Система 1. РИС-н в ряду. С экрана, τi = 0,5 мин (локальное время аппарата).
CSTR_SERIES = {
    1: [(0.0, 0.0000, 0.0000, 0.0000), (0.5, 0.2308, 0.1923, 0.0385)],
    2: [(0.0, 0.2308, 0.1923, 0.0385), (0.5, 0.4083, 0.3082, 0.1001)],
    3: [(0.0, 0.4083, 0.3082, 0.1001), (0.5, 0.5448, 0.3706, 0.1742)],
    4: [(0.0, 0.5448, 0.3706, 0.1742), (0.5, 0.6498, 0.3964, 0.2535)],
}
# Система 2. РИС-н параллельно. τi = Vi/vi = 25/12,5 = 2 мин. Как единичный РИС-н, л/р №3.
CSTR_PAR = [(0.0, 0.0000, 0.0000, 0.0000), (2.0, 0.5454, 0.3030, 0.2424)]


def xyz_from_c(ca, cr, cs):
    return 1.0 - ca, cr, cs


def pfr_step(tau, ca_in, cr_in, cs_in):
    if tau == 0:
        return ca_in, cr_in, cs_in
    ca = ca_in * math.exp(-K1 * tau)
    cr = cr_in * math.exp(-K2 * tau) + (K1 / (K2 - K1)) * ca_in * (
        math.exp(-K1 * tau) - math.exp(-K2 * tau)
    )
    cs = ca_in + cr_in + cs_in - ca - cr
    return ca, cr, cs


def pfr_ends(tau_max, ca_in=1.0, cr_in=0.0, cs_in=0.0):
    """Две строки как у РИС-н: вход τ=0 и выход τ=τi."""
    x0, y0, z0 = xyz_from_c(ca_in, cr_in, cs_in)
    ca, cr, cs = pfr_step(tau_max, ca_in, cr_in, cs_in)
    x, y, z = xyz_from_c(ca, cr, cs)
    rows = [
        (0.0, round(x0, 4), round(y0, 4), round(z0, 4)),
        (tau_max, round(x, 4), round(y, 4), round(z, 4)),
    ]
    return rows, ca, cr, cs


def rec(tau, x, y, z):
    return (tau, x, CA0 * (1.0 - x), y, CA0 * y, z, CA0 * z)


def comma(v, nd):
    return f"{v:.{nd}f}".replace(".", ",")


def set_run(run, size=14, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=14, bold=False, color=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, bold=False, size=11, fill=None, color=None, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
    )
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    set_cell_border(cell)
    if fill:
        shade(cell, fill)


def add_table(doc, rows):
    # Порядок колонок — как в бланке задания.
    headers = ["Тау, мин", "X", "CA, моль/л", "Y", "CR, моль/л", "Z", "CS, моль/л"]
    tbl = doc.add_table(rows=1 + len(rows), cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for j, h in enumerate(headers):
        fill_cell(tbl.rows[0].cells[j], h, bold=True, size=11, fill="1F4E79", color=(255, 255, 255))
    for i, raw in enumerate(rows):
        tau, x, ca, y, cr, z, cs = rec(*raw)
        vals = [
            comma(tau, 2),
            comma(x, 4),
            comma(ca, 2),
            comma(y, 4),
            comma(cr, 2),
            comma(z, 4),
            comma(cs, 2),
        ]
        last = i == len(rows) - 1
        fill = "FFF2CC" if last else "FFFFFF"
        for j, v in enumerate(vals):
            fill_cell(tbl.rows[i + 1].cells[j], v, bold=last, size=11, fill=fill)
    doc.add_paragraph()
    return tbl


def build():
    ca_in, cr_in, cs_in = 1.0, 0.0, 0.0
    pfr_series = []
    for _ in range(4):
        rows, ca_in, cr_in, cs_in = pfr_ends(0.5, ca_in, cr_in, cs_in)
        pfr_series.append(rows)
    pfr_par, *_ = pfr_ends(2.0, 1.0, 0.0, 0.0)

    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(1.8)
        sec.right_margin = Cm(1.5)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)

    add_p(doc, "Лабораторная работа «Реакторные системы». Изотермические процессы", size=16, bold=True, space_after=2)
    add_p(doc, "Вариант 2. n1 = n2 = 1.  CA0 = 30 моль/л, V = 100 л, v0 = 50 л/мин", size=14, bold=True, space_after=6)
    add_p(doc, "Пункт 2. Таблица с результатами моделирования для каждого реактора каждой системы", size=14, bold=True, space_after=8)
    add_p(
        doc,
        "По методичке: отдельная таблица на каждый аппарат. Колонки как в бланке. "
        "Тау — время пребывания в моделируемом реакторе (не сумма по системе). "
        "У РИС-н на экране две строки: вход τ = 0 и выход τ = τi. У РИВ — то же: вход и выход аппарата. "
        "CA = 30·(1−X), CR = 30·Y, CS = 30·Z. Жёлтая строка — выход реактора.",
        size=12,
        space_after=8,
    )

    add_p(
        doc,
        "Система 1. Четыре последовательно соединённых РИС-н одинакового объёма",
        size=14,
        bold=True,
        color=(31, 78, 121),
    )
    add_p(doc, "Vi = 25 л, vi = 50 л/мин, τi = 0,5 мин. Данные с экрана программы.", size=12)
    for i in range(1, 5):
        add_p(doc, f"Реактор {i}", size=13, bold=True, color=(46, 117, 182), space_before=4, space_after=4)
        add_table(doc, CSTR_SERIES[i])

    add_p(
        doc,
        "Система 2. Четыре параллельно соединённых РИС-н одинакового объёма, одинаковая нагрузка",
        size=14,
        bold=True,
        color=(31, 78, 121),
        space_before=8,
    )
    add_p(doc, "Vi = 25 л, vi = 12,5 л/мин, τi = 2 мин. Все четыре аппарата одинаковы (как единичный РИС-н, л/р №3).", size=12)
    for i in range(1, 5):
        add_p(doc, f"Реактор {i}", size=13, bold=True, color=(46, 117, 182), space_before=4, space_after=4)
        add_table(doc, CSTR_PAR)

    add_p(
        doc,
        "Система 3. Четыре последовательно соединённых РИВ одинакового объёма",
        size=14,
        bold=True,
        color=(31, 78, 121),
        space_before=8,
    )
    add_p(doc, "Vi = 25 л, vi = 50 л/мин, τi = 0,5 мин. Вход следующего = выход предыдущего.", size=12)
    for i, rows in enumerate(pfr_series, 1):
        add_p(doc, f"Реактор {i}", size=13, bold=True, color=(46, 117, 182), space_before=4, space_after=4)
        add_table(doc, rows)

    add_p(
        doc,
        "Система 4. Четыре параллельно соединённых РИВ одинакового объёма, одинаковая нагрузка",
        size=14,
        bold=True,
        color=(31, 78, 121),
        space_before=8,
    )
    add_p(doc, "Vi = 25 л, vi = 12,5 л/мин, τi = 2 мин. Все четыре аппарата одинаковы (как единичный РИВ, л/р №3).", size=12)
    for i in range(1, 5):
        add_p(doc, f"Реактор {i}", size=13, bold=True, color=(46, 117, 182), space_before=4, space_after=4)
        add_table(doc, pfr_par)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cyr = OUT_DIR / "Пункт2_таблицы_моделирования.docx"
    lat = OUT_DIR / "Punkt2_modeling_tables.docx"
    doc.save(cyr)
    shutil.copyfile(cyr, lat)
    print("saved", cyr)
    print("pfr series last", pfr_series[-1][-1])
    print("pfr par last", pfr_par[-1])


if __name__ == "__main__":
    build()
