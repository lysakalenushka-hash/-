#!/usr/bin/env python3
"""Контрольная ОХТ, вариант 19: задачи 1–2 по семинару РХТУ (Бесков / Царёва)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path("/workspace/oxt_var19")
LAT = OUT_DIR / "OXT_KR_variant_19.docx"
CYR = OUT_DIR / "ОХТ_КР_вариант_19.docx"


def set_run(run, size=14, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def P(doc, text, size=14, bold=False, italic=False, color=None, after=4, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        set_run(p.add_run(part), size=size, bold=bold, italic=italic, color=color)
    return p


def H1(doc, text):
    P(doc, text, size=16, bold=True, color=(31, 78, 121), before=14, after=8)


def given_find(doc, given, find):
    P(doc, "Дано.", size=14, bold=True, before=2, after=2)
    P(doc, given, size=14, after=4)
    P(doc, "Найти.", size=14, bold=True, after=2)
    P(doc, find, size=14, after=4)
    P(doc, "Решение.", size=14, bold=True, after=4)


def Ans(doc, text):
    P(doc, "Ответ: " + text, size=14, bold=True, color=(0, 100, 0), before=6, after=12)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(1.6)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(t.add_run("РХТУ им. Д. И. Менделеева"), size=14, bold=True)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(t.add_run("Общая химическая технология"), size=16, bold=True)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(
        t.add_run("Контрольная работа. Вариант 19"),
        size=16,
        bold=True,
    )

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(
        t.add_run("Физико-химические основы химико-технологических процессов"),
        size=13,
        italic=True,
    )

    P(
        doc,
        "Задачи 1 и 2 решены по конспекту семинара (стехиометрия, выход, селективность; "
        "равновесие Kc = k1/k−1) и по Бескову. Объём жидкой фазы постоянен: концентрации "
        "в кмоль/м³ можно подставлять вместо количеств вещества. "
        "R = 8,314 Дж/(моль·К).",
        size=13,
        italic=True,
        before=10,
        after=10,
    )

    # ------------------------------------------------------------------
    H1(doc, "Задача 1. Стехиометрия сложного процесса")
    given_find(
        doc,
        "Жидкофазный процесс:\n"
        "    (1)  A + 3B → C + S\n"
        "    (2)  2C + A → D + 2S\n"
        "    (3)  3A + 2B → 4P     (целевая)\n"
        "C_A0 = 36 кмоль/м³;  C_B0 = 24 кмоль/м³;\n"
        "C_A = 12;  C_S = 14;  C_P = 5 кмоль/м³;\n"
        "C_D0 = C_S0 = C_P0 = 0  (C и D в начале тоже отсутствуют).",
        "x_B,  E_P,  S_P(A).",
    )
    P(
        doc,
        "Обозначим ξ1, ξ2, ξ3 — степени протекания реакций (1)–(3), кмоль/м³. "
        "Тогда (V = const):\n"
        "    C_A = C_A0 − ξ1 − ξ2 − 3ξ3,\n"
        "    C_B = C_B0 − 3ξ1 − 2ξ3,\n"
        "    C_C = ξ1 − 2ξ2,\n"
        "    C_D = ξ2,\n"
        "    C_S = ξ1 + 2ξ2,\n"
        "    C_P = 4ξ3.\n"
        "Из C_P = 5:\n"
        "    ξ3 = 5 / 4 = 1,25 кмоль/м³.\n"
        "Из C_A = 12:\n"
        "    ξ1 + ξ2 + 3·1,25 = 36 − 12 = 24  ⇒  ξ1 + ξ2 = 20,25.     (а)\n"
        "Из C_S = 14:\n"
        "    ξ1 + 2ξ2 = 14.                                            (б)\n"
        "(б) − (а):  ξ2 = 14 − 20,25 = −6,25 кмоль/м³,\n"
        "            ξ1 = 20,25 − (−6,25) = 26,50 кмоль/м³.\n"
        "Отрицательное ξ2 означает, что тройка чисел (C_A, C_S, C_P) линейно несовместима "
        "с необратимой схемой (1)–(3) при неотрицательных концентрациях C и D. "
        "Ниже считаем показатели, которые от ξ2 не зависят, — они однозначны.",
    )
    P(
        doc,
        "Целевой продукт P образуется только в реакции (3). Поэтому расход A на P "
        "и выход P фиксируются одним C_P (формулы семинара, Бесков):\n"
        "    ΔC_A→P = (ν_A(3) / ν_P) · C_P = (3/4) · 5 = 3,75 кмоль/м³.\n"
        "Прореагировало A:\n"
        "    ΔC_A = C_A0 − C_A = 36 − 12 = 24 кмоль/м³,\n"
        "    x_A = ΔC_A / C_A0 = 24 / 36 = 0,667.\n"
        "Интегральная селективность P по исходному A:\n"
        "    S_P(A) = ΔC_A→P / ΔC_A = 3,75 / 24 = 0,156.\n"
        "Выход P (доля A, ушедшего в целевой продукт, от загруженного A; "
        "эквивалентно C_P / C_P^max при x_A = 1 и S_P = 1):\n"
        "    C_P^max = (4/3) C_A0 = 48 кмоль/м³,\n"
        "    E_P = C_P / C_P^max = 5 / 48 = 0,104.\n"
        "Проверка связи семинара E = x · S:\n"
        "    E_P = x_A · S_P(A) = 0,667 · 0,156 = 0,104.",
    )
    P(
        doc,
        "Степень превращения B. По определению x_B = ΔC_B / C_B0, "
        "ΔC_B = 3ξ1 + 2ξ3. Подстановка формального решения системы даёт "
        "x_B > 1 (нефизично) из‑за ξ2 < 0.\n"
        "Поэтому x_B из бланка однозначно не восстанавливается: для B не хватает "
        "согласованного C_D (или C_C). В типовой задаче кафедры 1.2 как раз задают C_D, "
        "а C_P находят из баланса.\n"
        "Если принять, что в бланке вместо C_P указан C_D = 5 кмоль/м³ "
        "(как в задаче 1.2), то ξ2 = 5, ξ1 = 14 − 10 = 4, ξ3 = (24 − 4 − 5)/3 = 5, "
        "C_P = 20, C_B = 24 − 12 − 10 = 2,\n"
        "    x_B = (24 − 2) / 24 = 0,917,\n"
        "    S_P(A) = 3·5 / 24 = 0,625,\n"
        "    E_P = 15 / 36 = 0,417,\n"
        "но тогда C_C = 4 − 10 = −6 < 0 — тоже нефизично.\n"
        "В ответ по тексту задания (дан C_P = 5) записываем однозначные E_P и S_P(A); "
        "x_B по согласованному балансу B в бланке не определяется.",
        size=13,
    )
    Ans(
        doc,
        "E_P = 0,104;   S_P(A) = 0,156.   "
        "x_B по данным бланка однозначно не находится (несовместны C_A, C_S и C_P).",
    )

    # ------------------------------------------------------------------
    H1(doc, "Задача 2. Равновесие A  ↔  2B")
    given_find(
        doc,
        "Жидкофазная обратимая реакция A ↔ 2B, t = 40 °C (T = 313,15 К).\n"
        "C_A0 = 3 кмоль/м³,  C_B0 = 0.\n"
        "    k1  = 5·10³ exp(−15000 / RT)   [мин⁻¹],\n"
        "    k−1 = 2·10³ exp(−27000 / RT)   [л/(моль·мин)].\n"
        "R = 8,314 Дж/(моль·К).  1 кмоль/м³ = 1 моль/л.",
        "Равновесный состав C_A*, C_B*.",
    )
    P(
        doc,
        "Прямая реакция 1-го порядка, обратная — 2-го (по двум молекулам B). "
        "Константа равновесия по концентрациям (семинар 2, Бесков):\n"
        "    K_c = k1 / k−1 = C_B*² / C_A*     [моль/л].\n"
        "Считаем константы при T = 313,15 К:\n"
        "    15000 / (R T) = 15000 / (8,314 · 313,15) = 5,761,\n"
        "    k1  = 5·10³ · e^(−5,761) = 5·10³ · 0,003147 = 15,73 мин⁻¹.\n"
        "    27000 / (R T) = 27000 / (8,314 · 313,15) = 10,370,\n"
        "    k−1 = 2·10³ · e^(−10,370) = 2·10³ · 3,134·10⁻⁵ = 0,0627 л/(моль·мин).\n"
        "    K_c = 15,73 / 0,0627 = 251,0 моль/л.",
    )
    P(
        doc,
        "Материальный баланс (C_B0 = 0):\n"
        "    C_A* = C_A0 − ξ* = 3 − ξ*,\n"
        "    C_B* = 2ξ* = 2 (3 − C_A*).\n"
        "Подстановка в K_c:\n"
        "    4 (3 − C_A*)² / C_A* = 251,0,\n"
        "    4 C_A*² − (24 + 251) C_A* + 36 = 0,\n"
        "    4 C_A*² − 275 C_A* + 36 = 0.\n"
        "    D = 275² − 4·4·36 = 75625 − 576 = 75049,\n"
        "    C_A* = [275 − √75049] / 8.\n"
        "Физический корень на (0; 3):\n"
        "    C_A* = 0,131 кмоль/м³,\n"
        "    C_B* = 2 (3 − 0,131) = 5,74 кмоль/м³.\n"
        "Проверка:  C_B*² / C_A* = 5,74² / 0,131 = 251 = K_c.\n"
        "Равновесная степень превращения A:  x_A* = 1 − 0,131/3 = 0,956.\n"
        "Большое K_c (энергия активации обратной реакции выше) смещает равновесие вправо, "
        "почти весь A превращается в B.",
    )
    Ans(doc, "C_A* = 0,131 кмоль/м³;   C_B* = 5,74 кмоль/м³.")

    P(
        doc,
        "Примечание. Вопросы 3 (структура ХП) и 4 (кинетические закономерности) в расчёт "
        "не входят — по условию пользователя решены только две задачи.",
        size=12,
        italic=True,
        before=8,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(LAT)
    shutil.copyfile(LAT, CYR)
    print("saved", LAT, LAT.stat().st_size)


if __name__ == "__main__":
    build()
