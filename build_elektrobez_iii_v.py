#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""РТН приложение 2: разделы III и V + сопоставление с анализом блог-инженера."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RTN = Path("Приложение_2_РТН.xlsx")
BLOG = Path("Voprosi-po-elektrobezopasnosti-2026.xlsx")
OUTPUT = Path("Электробезопасность_РТН_III_и_V.xlsx")

MAX_ANS = 6
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
ALIGN = Alignment(wrap_text=True, vertical="top")
FONT = Font(name="Times New Roman", size=11)
FONT_B = Font(name="Times New Roman", size=11, bold=True)
FONT_H = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="1F4E79")
FILL_NPA = PatternFill("solid", fgColor="D6EAF8")
FILLS = {
    "Без изменений": PatternFill("solid", fgColor="C6EFCE"),
    "Изменена формулировка": PatternFill("solid", fgColor="FFF2CC"),
    "Добавлен": PatternFill("solid", fgColor="C9DAF8"),
    "Удалён": PatternFill("solid", fgColor="F4C7C3"),
}


def norm(s: str) -> str:
    s = (s or "").lower().replace("\xa0", " ").replace("ё", "е")
    s = re.sub(r"[«»“”]", '"', s)
    return re.sub(r"\s+", " ", s).strip()


def parse_rtn(ws, has_meta: bool = False):
    npa = None
    questions = []
    cur = None
    seq = 0
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a == "№":
            continue
        is_num = isinstance(a, int) or (isinstance(a, str) and str(a).strip().isdigit())
        if a and not is_num:
            npa = str(a).strip()
            continue
        if is_num and b:
            if cur:
                questions.append(cur)
            seq += 1
            cur = {
                "seq": seq,
                "local": int(str(a).strip()),
                "npa": npa or "",
                "q": str(b).strip(),
                "ans": [],
                "до 1000 В": "",
                "до и выше 1000 В": "",
                "II": "",
                "III": "",
                "IV": "",
                "V": "",
            }
            if has_meta:
                def plus(c):
                    return "+" if ws.cell(r, c).value == "+" else ""

                cur["до 1000 В"] = plus(3)
                cur["до и выше 1000 В"] = plus(4)
                cur["II"] = plus(5)
                cur["III"] = plus(6)
                cur["IV"] = plus(7)
                cur["V"] = plus(8)
            continue
        if cur and b and (a is None or a == ""):
            cur["ans"].append(str(b).strip())
    if cur:
        questions.append(cur)
    return questions


def load_blog(path: Path, section: str):
    wb = load_workbook(path, data_only=True)
    ws = wb["Сравнение"]
    rows = []
    by_new = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value != section:
            continue
        rec = {
            "status": ws.cell(r, 2).value or "",
            "old_n": ws.cell(r, 3).value,
            "old_q": ws.cell(r, 4).value or "",
            "new_n": ws.cell(r, 5).value,
            "new_q": ws.cell(r, 6).value or "",
            "old_a": ws.cell(r, 7).value or "",
            "new_a": ws.cell(r, 8).value or "",
            "ch_a": ws.cell(r, 9).value or "",
            "sim": ws.cell(r, 11).value,
        }
        rows.append(rec)
        nq = norm(rec["new_q"])
        if nq and rec["status"] != "Удалён":
            by_new.setdefault(nq, rec)
    return rows, by_new


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = FONT_H
        cell.fill = FILL_H
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = f"A{row + 1}"


def write_row(ws, r, values, fill=None, bold=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(r, c, v)
        cell.font = FONT_B if bold else FONT
        cell.alignment = ALIGN
        cell.border = THIN
        if fill is not None:
            cell.fill = fill


def join_answers(ans):
    return "\n".join(f"{i}. {t}" for i, t in enumerate(ans, 1))


def headers(has_meta: bool):
    h = [
        "№ РТН",
        "№ в блоке НПА",
        "НПА / блок",
        "Статус (блог-инженер)",
        "Вопрос (РТН, 2026)",
        *[f"Ответ {i}" for i in range(1, MAX_ANS + 1)],
        "Варианты ответов (списком)",
        "№ старого",
        "Старый вопрос",
        "Изменение вариантов",
        "Сходство, %",
        "Старые варианты ответов",
    ]
    if has_meta:
        h[5:5] = ["до 1000 В", "до и выше 1000 В", "II", "III", "IV", "V"]
    return h


def meta_vals(q, has_meta):
    if not has_meta:
        return []
    return [q["до 1000 В"], q["до и выше 1000 В"], q["II"], q["III"], q["IV"], q["V"]]


def question_row(q, rec, has_meta):
    rec = rec or {}
    ans = q["ans"] + [""] * (MAX_ANS - len(q["ans"]))
    return [
        q["seq"],
        q["local"],
        q["npa"],
        rec.get("status") or "Нет в анализе",
        q["q"],
        *meta_vals(q, has_meta),
        *ans[:MAX_ANS],
        join_answers(q["ans"]),
        rec.get("old_n"),
        rec.get("old_q") or None,
        rec.get("ch_a") or None,
        rec.get("sim"),
        rec.get("old_a") or None,
    ]


def write_questions(ws, qs, by_new, has_meta, only_changed=False):
    h = headers(has_meta)
    write_row(ws, 1, h, fill=FILL_H, bold=True)
    style_header(ws, 1, len(h))
    r = 2
    n_written = 0
    for q in qs:
        rec = by_new.get(norm(q["q"]))
        st = (rec or {}).get("status") or "Нет в анализе"
        if only_changed and st == "Без изменений":
            continue
        fill = FILLS.get(st)
        write_row(ws, r, question_row(q, rec, has_meta), fill=fill)
        ws.row_dimensions[r].height = 48
        r += 1
        n_written += 1
    apply_widths(ws, has_meta)
    last = max(r - 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(h))}{last}"
    return n_written


def write_deleted(ws, blog_rows):
    h = ["№ старого", "Старый вопрос", "Старые варианты ответов", "Статус"]
    write_row(ws, 1, h, fill=FILL_H, bold=True)
    style_header(ws, 1, len(h))
    r = 2
    for rec in blog_rows:
        if rec["status"] != "Удалён":
            continue
        write_row(
            ws,
            r,
            [rec["old_n"], rec["old_q"], rec["old_a"], rec["status"]],
            fill=FILLS["Удалён"],
        )
        ws.row_dimensions[r].height = 48
        r += 1
    last = max(r - 1, 1)
    ws.auto_filter.ref = f"A1:D{last}"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 16
    return r - 2


def apply_widths(ws, has_meta):
    widths = {
        "A": 10,
        "B": 12,
        "C": 42,
        "D": 22,
        "E": 55,
    }
    start = 6
    if has_meta:
        for i, w in enumerate([12, 16, 6, 6, 6, 6], start):
            ws.column_dimensions[get_column_letter(i)].width = w
        start = 12
    for i in range(MAX_ANS):
        ws.column_dimensions[get_column_letter(start + i)].width = 28
    rest = start + MAX_ANS
    extra = [40, 12, 50, 28, 12, 40]
    for i, w in enumerate(extra):
        ws.column_dimensions[get_column_letter(rest + i)].width = w
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def npa_table(qs):
    c = Counter(q["npa"] for q in qs)
    return list(c.items())


def write_summary(ws, iii, v, blog_iii, blog_v, by_iii, by_v):
    ws["A1"] = "Электробезопасность: вопросы РТН (приложение 2) — разделы III, затем V"
    ws["A1"].font = Font(name="Times New Roman", size=16, bold=True, color="1F4E79")
    ws.merge_cells("A1:G1")
    notes = [
        "Источник «РТН» — официальное приложение 2 с сайта Ростехнадзора (файл Приложение_2_РТН.xlsx).",
        "Источник «блог-инженер» — сравнение старой и новой редакции с сайта блог-инженера (Voprosi-po-elektrobezopasnosti-2026.xlsx).",
        "По РТН взяты только раздел III (потребители тепловой энергии), затем раздел V (потребители электрической энергии).",
        "Нумерация «№ РТН» сквозная внутри раздела. «№ в блоке НПА» — как в исходном приложении РТН (сбрасывается на каждом блоке).",
        "Правильный ответ в файле РТН не отмечен — как и в анализе блог-инженера.",
        "Цвета статуса: зелёный — без изменений, жёлтый — изменена формулировка, синий — добавлен, красный — удалён.",
    ]
    for i, t in enumerate(notes, 3):
        ws.cell(i, 1, t).font = FONT
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
        ws.row_dimensions[i].height = 20

    write_row(
        ws,
        10,
        ["Раздел", "Вопросов РТН", "Без изменений", "Изменена формулировка", "Добавлен", "Удалён (только в старой редакции)", "Не найден в анализе"],
        fill=FILL_H,
        bold=True,
    )
    style_header(ws, 10, 7)

    def stats(qs, by_new, blog_rows):
        st = Counter()
        for q in qs:
            rec = by_new.get(norm(q["q"]))
            st[(rec or {}).get("status") or "Нет в анализе"] += 1
        deleted = sum(1 for x in blog_rows if x["status"] == "Удалён")
        return len(qs), st, deleted

    row = 11
    for title, qs, by_new, blog_rows in (
        ("III — потребители тепловой энергии", iii, by_iii, blog_iii),
        ("V — потребители электрической энергии", v, by_v, blog_v),
    ):
        n, st, deleted = stats(qs, by_new, blog_rows)
        write_row(
            ws,
            row,
            [
                title,
                n,
                st.get("Без изменений", 0),
                st.get("Изменена формулировка", 0),
                st.get("Добавлен", 0),
                deleted,
                st.get("Нет в анализе", 0),
            ],
        )
        row += 1

    write_row(ws, 14, ["Раздел III — блоки НПА", "Вопросов"], fill=FILL_NPA, bold=True)
    r = 15
    for name, cnt in npa_table(iii):
        write_row(ws, r, [name, cnt])
        r += 1
    r += 1
    write_row(ws, r, ["Раздел V — блоки НПА", "Вопросов", "из них с отметкой группы III", "из них с отметкой группы V"], fill=FILL_NPA, bold=True)
    r += 1
    start = r
    npa_c = Counter()
    g3 = Counter()
    g5 = Counter()
    for q in v:
        npa_c[q["npa"]] += 1
        if q["III"] == "+":
            g3[q["npa"]] += 1
        if q["V"] == "+":
            g5[q["npa"]] += 1
    for name, cnt in npa_c.items():
        write_row(ws, r, [name, cnt, g3[name], g5[name]])
        r += 1

    ws.column_dimensions["A"].width = 88
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 22
    ws.row_dimensions[1].height = 28


def build():
    rtn_wb = load_workbook(RTN, data_only=True)
    iii = parse_rtn(rtn_wb["III"])
    v = parse_rtn(rtn_wb["V"], has_meta=True)
    blog_iii, by_iii = load_blog(BLOG, "III")
    blog_v, by_v = load_blog(BLOG, "V")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Сводка"
    write_summary(ws0, iii, v, blog_iii, blog_v, by_iii, by_v)

    ws = wb.create_sheet("III_РТН_все")
    write_questions(ws, iii, by_iii, has_meta=False)
    ws = wb.create_sheet("III_изменения")
    write_questions(ws, iii, by_iii, has_meta=False, only_changed=True)
    ws = wb.create_sheet("III_удалённые")
    write_deleted(ws, blog_iii)

    ws = wb.create_sheet("V_РТН_все")
    write_questions(ws, v, by_v, has_meta=True)
    ws = wb.create_sheet("V_изменения")
    write_questions(ws, v, by_v, has_meta=True, only_changed=True)
    ws = wb.create_sheet("V_удалённые")
    write_deleted(ws, blog_v)

    wb.save(OUTPUT)
    print(f"Saved {OUTPUT}: III={len(iii)}, V={len(v)}")


if __name__ == "__main__":
    build()
