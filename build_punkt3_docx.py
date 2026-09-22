#!/usr/bin/env python3
"""Пункт 3. Производительность ΠR и сравнение с л/р №3 (n1 = n2)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT_DIR = Path("/workspace/reaktory_sistemy_variant2")
CA0 = 30.0
V0 = 50.0  # л/мин
PI = V0 * 60 / 1000.0  # 3 м³/ч; ΠR = 3·CR = 90·Y, кмоль/ч

# Y выхода: система 1 с экрана; 2–4 как единичные л/р №3, n1 = n2 = 1
ROWS = [
    ("Единичный РИС-н, л/р №3", 0.5454, 0.3030, "база для системы 2"),
    ("Единичный РИВ, л/р №3", 0.6988, 0.4444, "база для систем 3 и 4"),
    ("1. 4 РИС-н последовательно", 0.6498, 0.3964, "выше РИС-н, ниже РИВ"),
    ("2. 4 РИС-н параллельно", 0.5454, 0.3030, "совпадает с единичным РИС-н"),
    ("3. 4 РИВ последовательно", 0.6988, 0.4444, "совпадает с единичным РИВ"),
    ("4. 4 РИВ параллельно", 0.6988, 0.4444, "совпадает с единичным РИВ"),
]


def comma(v, nd):
    return f"{v:.{nd}f}".replace(".", ",")


def set_run(run, size=14, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=14, bold=False, color=None, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, bold=False, size=11, fill=None, color=None, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    set_cell_border(cell)
    if fill:
        shade(cell, fill)


def build():
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(1.5)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)

    add_p(doc, "Лабораторная работа «Реакторные системы». Вариант 2", size=16, bold=True, space_after=2)
    add_p(
        doc,
        "Пункт 3. Производительность систем по целевому продукту R и сравнение с л/р №3 (n1 = n2 = 1)",
        size=14,
        bold=True,
        space_after=8,
    )

    add_p(doc, "Формула", size=14, bold=True, color=(31, 78, 121))
    add_p(
        doc,
        "ΠR = v0 · CR_вых. Расход через систему v0 = 50 л/мин = 3 м³/ч. "
        "CR в моль/л численно равен кмоль/м³, поэтому ΠR = 3 · CR, кмоль/ч. "
        "CR = 30 · Y, значит ΠR = 90 · Y. Y берётся с выхода системы (последний аппарат или любой из параллельных).",
        size=14,
    )

    add_p(doc, "Подстановка для каждой схемы", size=14, bold=True, color=(31, 78, 121), space_before=8)

    calcs = [
        ("Единичный РИС-н, л/р №3", 0.3030),
        ("Единичный РИВ, л/р №3", 0.4444),
        ("Система 1. 4 РИС-н последовательно (Y с выхода реактора 4)", 0.3964),
        ("Система 2. 4 РИС-н параллельно", 0.3030),
        ("Система 3. 4 РИВ последовательно (Y с выхода реактора 4)", 0.4444),
        ("Система 4. 4 РИВ параллельно", 0.4444),
    ]
    for title, y in calcs:
        cr = CA0 * y
        pir = 90.0 * y
        add_p(doc, title, size=13, bold=True, color=(46, 117, 182), space_before=6, space_after=3)
        add_p(
            doc,
            f"Y = {comma(y, 4)};  CR = 30 · {comma(y, 4)} = {comma(cr, 2)} моль/л;\n"
            f"ΠR = 90 · {comma(y, 4)} = {comma(pir, 2)} кмоль/ч.",
            size=14,
            space_after=4,
        )

    add_p(doc, "Сводная таблица", size=14, bold=True, color=(31, 78, 121), space_before=8)
    headers = ["Схема", "Y вых.", "CR, моль/л", "ΠR, кмоль/ч", "Сравнение с л/р №3"]
    tbl = doc.add_table(rows=1 + len(ROWS), cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for j, h in enumerate(headers):
        fill_cell(tbl.rows[0].cells[j], h, bold=True, size=11, fill="1F4E79", color=(255, 255, 255))
    for i, (name, _x, y, note) in enumerate(ROWS):
        cr = CA0 * y
        pir = 90.0 * y
        fill = "E2EFDA" if "РИВ" in name else ("FFF2CC" if i >= 2 else "D6EAF8")
        vals = [name, comma(y, 4), comma(cr, 2), comma(pir, 2), note]
        aligns = ["left", "center", "center", "center", "left"]
        for j, (v, a) in enumerate(zip(vals, aligns)):
            fill_cell(tbl.rows[i + 1].cells[j], v, bold=(j == 0), size=11, fill=fill, align=a)

    add_p(doc, "Почему так", size=14, bold=True, color=(31, 78, 121), space_before=12)
    add_p(
        doc,
        "Система 2 совпадает с единичным РИС-н (27,27 кмоль/ч): четыре одинаковых аппарата с равной нагрузкой "
        "работают как один РИС-н того же времени пребывания, что в л/р №3. "
        "Системы 3 и 4 совпадают с единичным РИВ (40,00 кмоль/ч): четыре РИВ в ряду — это один РИВ суммарного объёма; "
        "четыре РИВ параллельно при равной нагрузке — снова тот же РИВ, что в л/р №3.",
        size=14,
    )
    add_p(
        doc,
        "Система 1 даёт 35,68 кмоль/ч — между 27,27 и 40,00. Каскад четырёх РИС-н сглаживает проскок "
        "и приближает режим к вытеснению, но полностью его не достигает, поэтому ΠR выше, чем у одного РИС-н, "
        "и ниже, чем у РИВ.",
        size=14,
    )
    add_p(
        doc,
        "Итог по ΠR: 4 РИВ (ряд или параллель, 40,00) > 4 РИС-н в ряду (35,68) > 4 РИС-н параллельно (27,27).",
        size=14,
        bold=True,
        space_before=6,
    )

    chart = OUT_DIR / "graf_proizvoditelnost.png"
    if chart.exists():
        add_p(doc, "Диаграмма ΠR", size=14, bold=True, color=(31, 78, 121), space_before=10)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(chart), width=Inches(6.2))

    cyr = OUT_DIR / "Пункт3_производительность.docx"
    lat = OUT_DIR / "Punkt3_productivity.docx"
    doc.save(cyr)
    shutil.copyfile(cyr, lat)
    print("saved", cyr)


if __name__ == "__main__":
    build()
