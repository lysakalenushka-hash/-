#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Отчёт лабораторной №1, вариант 1 — схема пользователя без правок, остальное по примеру."""

from __future__ import annotations

import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mp
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
etree.register_namespace("m", M_NS)

OUT_DIR = Path("лаба1_вариант1")
DOCX = Path("Лабораторная_работа_1_гидравлика_вариант_1.docx")
USER_SCHEME = Path(
    "/home/ubuntu/.cursor/projects/workspace/assets/2c8b6ada-6b5a-4431-a349-7adf59a71ca4.png"
)


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_p(doc, text, *, size=14, bold=False, italic=False, align="left", space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def _mr(text, *, sub=False):
    r = etree.Element("{%s}r" % M_NS)
    rpr = etree.SubElement(r, "{%s}rPr" % W_NS)
    fonts = etree.SubElement(rpr, "{%s}rFonts" % W_NS)
    fonts.set(qn("w:ascii"), "Cambria Math")
    fonts.set(qn("w:hAnsi"), "Cambria Math")
    etree.SubElement(rpr, "{%s}sz" % W_NS).set(qn("w:val"), "28")
    etree.SubElement(rpr, "{%s}szCs" % W_NS).set(qn("w:val"), "28")
    if sub:
        etree.SubElement(rpr, "{%s}vertAlign" % W_NS).set(qn("w:val"), "subscript")
    mpr = etree.SubElement(r, "{%s}rPr" % M_NS)
    etree.SubElement(mpr, "{%s}sty" % M_NS).set(qn("m:val"), "p")
    t = etree.SubElement(r, "{%s}t" % M_NS)
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return r


def omml_sqrt_abs(left: str, right: str):
    omath = etree.Element("{%s}oMath" % M_NS)
    omath.append(_mr("√"))
    d = etree.SubElement(omath, "{%s}d" % M_NS)
    dpr = etree.SubElement(d, "{%s}dPr" % M_NS)
    etree.SubElement(dpr, "{%s}begChr" % M_NS).set(qn("m:val"), "|")
    etree.SubElement(dpr, "{%s}endChr" % M_NS).set(qn("m:val"), "|")
    e = etree.SubElement(d, "{%s}e" % M_NS)
    e.append(_mr("Р"))
    e.append(_mr(left, sub=True))
    e.append(_mr(" – Р"))
    e.append(_mr(right, sub=True))
    return omath


def omml_pn():
    omath = etree.Element("{%s}oMath" % M_NS)
    ssup = etree.SubElement(omath, "{%s}sSup" % M_NS)
    e = etree.SubElement(ssup, "{%s}e" % M_NS)
    e.append(_mr("Р"))
    up = etree.SubElement(ssup, "{%s}sup" % M_NS)
    up.append(_mr("N"))
    return omath


def append_runs(p, parts):
    for part in parts:
        if isinstance(part, str):
            run = p.add_run(part)
            set_run_font(run)
        else:
            p._p.append(part)


def draw_matrix(path: Path) -> None:
    # V1–V7, P6–P8 жидкость, P9–P11 газ, H1–H3
    cols = [
        "V1", "V2", "V3", "V4", "V5", "V6", "V7",
        "P6", "P7", "P8", "P9", "P10", "P11",
        "H1", "H2", "H3",
    ]
    rows = ["1", "2", "3", "4", "5", "6", "7", "8'", "9'", "10'", "11", "12", "13", "14", "15", "16"]
    # определяемые: V1–V7, H1, H2, H3, P6, P9, P7, P10, P8, P11
    det = [0, 1, 2, 3, 4, 5, 6, 13, 14, 15, 7, 10, 8, 11, 9, 12]
    used = {
        0: [0, 7],
        1: [1, 7],
        2: [2, 7],
        3: [3, 7, 8],
        4: [4, 8, 9],
        5: [5, 8],
        6: [6, 9],
        7: [0, 1, 2, 3, 13],
        8: [3, 4, 5, 14],
        9: [4, 6, 15],
        10: [7, 10, 13],
        11: [10, 13],
        12: [8, 11, 14],
        13: [11, 14],
        14: [9, 12, 15],
        15: [12, 15],
    }
    # газ раньше Паскаля, затем V, затем балансы
    step_no = [7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 2, 1, 4, 3, 6, 5]

    fig, ax = plt.subplots(figsize=(13.4, 7.8))
    ax.set_xlim(-1.7, len(cols) + 2.3)
    ax.set_ylim(-1.0, len(rows) + 2.2)
    ax.axis("off")
    ax.set_title("Информационная матрица системы уравнений", fontsize=13)
    ax.text(-0.9, len(rows) + 0.55, "№ ур.", ha="center", va="center", fontsize=8, weight="bold")
    ax.text(len(cols) + 1.15, len(rows) + 0.55, "№ шага", ha="center", va="center", fontsize=8, weight="bold")
    for j, c in enumerate(cols):
        ax.text(j + 0.5, len(rows) + 0.55, c, ha="center", va="center", fontsize=8, weight="bold")
        ax.plot([j, j], [0, len(rows) + 1.05], color="#cccccc", lw=0.45)
    ax.plot([0, len(cols)], [len(rows) + 1.05, len(rows) + 1.05], "k", lw=0.8)
    ax.plot([0, len(cols)], [0, 0], "k", lw=0.8)

    for i, num in enumerate(rows):
        y = len(rows) - i - 0.5
        ax.text(-0.9, y, num, ha="center", va="center", fontsize=8)
        for j in used[i]:
            ax.add_patch(plt.Circle((j + 0.5, y), 0.15, fill=False, lw=1.05, ec="k"))
        jdet = det[i]
        ax.plot(
            [jdet + 0.5, jdet + 0.72, jdet + 0.5, jdet + 0.28, jdet + 0.5],
            [y + 0.20, y, y - 0.20, y, y + 0.20],
            "k",
            lw=1.2,
        )
        ax.text(len(cols) + 1.15, y, str(step_no[i]), ha="center", va="center", fontsize=8)

    ax.text(
        len(cols) / 2,
        -0.5,
        "○ — переменная входит в уравнение;  ◇ — определяемая переменная.  "
        "P1–P5, ki, Si, H1G–H3G, PN, ρ, g — известны.",
        ha="center",
        fontsize=8,
    )
    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def _rect(ax, x, y, w, h, text, eq=None, fs=9):
    ax.add_patch(mp.Rectangle((x, y), w, h, fill=False, lw=1.15, ec="k"))
    ax.text(x + w / 2, y + h / 2 - (0.06 if eq else 0), text, ha="center", va="center", fontsize=fs)
    if eq is not None:
        ax.text(x + 0.07, y + h - 0.08, str(eq), ha="left", va="top", fontsize=7)


def _arrow(ax, x0, y0, x1, y1):
    ax.annotate(
        "",
        xy=(x1, y1),
        xytext=(x0, y0),
        arrowprops=dict(arrowstyle="->", lw=1.05, color="k"),
    )


def draw_flow(path: Path) -> None:
    """Слева направо, как в примерной лабе: три ёмкости, входы столбиком."""
    fig, ax = plt.subplots(figsize=(16.8, 8.4))
    ax.set_xlim(0.0, 16.8)
    ax.set_ylim(0.10, 8.70)
    ax.axis("off")

    y1, y2, y3 = 6.40, 3.55, 0.85
    bw, bh = 1.18, 0.72
    yc = (y1 + y3) / 2 + bh / 2

    ax.add_patch(mp.Ellipse((0.82, yc), 1.32, 0.82, fill=False, lw=1.2))
    ax.text(0.82, yc, "СТАРТ", ha="center", va="center", fontsize=10)

    _rect(ax, 1.65, 1.55, 1.72, 4.70, "")
    ax.text(
        2.51,
        yc,
        "ВВОД\nP1–P5\nk1–k7\nPN, ρ, g\nH1(o), H2(o), H3(o)\nS1, S2, S3\nH1G, H2G, H3G",
        ha="center",
        va="center",
        fontsize=8,
    )
    _arrow(ax, 1.48, yc, 1.65, yc)

    xh0, xg, xp, xv_in, xv_c, xv, xht, xout = 3.55, 4.85, 6.30, 7.70, 9.10, 10.45, 11.90, 13.55

    for y, lab in ((y1, "H1(o)"), (y2, "H2(o)"), (y3, "H3(o)")):
        _rect(ax, xh0, y, 1.05, bh, lab, fs=9)
        _arrow(ax, 3.37, y + bh / 2, xh0, y + bh / 2)

    for y, name, eq in ((y1, "P9", "12"), (y2, "P10", "14"), (y3, "P11", "16")):
        _rect(ax, xg, y, bw, bh, name, eq=eq)
        _arrow(ax, xh0 + 1.05, y + bh / 2, xg, y + bh / 2)

    for y, name, eq in ((y1, "P6", "11"), (y2, "P7", "13"), (y3, "P8", "15")):
        _rect(ax, xp, y, bw, bh, name, eq=eq)
        _arrow(ax, xg + bw, y + bh / 2, xp, y + bh / 2)

    yv1, yv2, yv3 = y1 + 1.05, y1, y1 - 1.05
    _rect(ax, xv_in, yv1, bw, bh, "V1", eq="1")
    _rect(ax, xv_in, yv2, bw, bh, "V2", eq="2")
    _rect(ax, xv_in, yv3, bw, bh, "V3", eq="3")
    for yv in (yv1, yv2, yv3):
        _arrow(ax, xp + bw, y1 + bh / 2, xv_in, yv + bh / 2)

    yv4 = (y1 + y2) / 2
    _rect(ax, xv_c, yv4, bw, bh, "V4", eq="4")
    _arrow(ax, xv_in + bw, y1 + 0.08, xv_c, yv4 + bh - 0.08)
    _arrow(ax, xp + bw, y2 + bh - 0.08, xv_c, yv4 + 0.08)

    _rect(ax, xv, y2 + 0.95, bw, bh, "V5", eq="5")
    _rect(ax, xv, y2 - 0.95, bw, bh, "V6", eq="6")
    _arrow(ax, xp + bw, y2 + bh / 2, xv, y2 + 0.95 + bh / 2)
    _arrow(ax, xp + bw, y2 + bh / 2, xv, y2 - 0.95 + bh / 2)
    _arrow(ax, xv_c + bw, yv4 + 0.10, xv, y2 + 0.95 + bh / 2)

    _rect(ax, xv, y3, bw, bh, "V7", eq="7")
    _arrow(ax, xp + bw, y3 + bh / 2, xv, y3 + bh / 2)

    _rect(ax, xht, y1, 1.42, bh, "H1(t(k))", eq="8'")
    _rect(ax, xht, y2, 1.42, bh, "H2(t(k))", eq="9'")
    _rect(ax, xht, y3, 1.42, bh, "H3(t(k))", eq="10'")
    for yv in (yv1, yv2, yv3):
        _arrow(ax, xv_in + bw, yv + bh / 2, xht, y1 + bh / 2)
    _arrow(ax, xv_c + bw, yv4 + bh / 2, xht, y1 + 0.10)
    _arrow(ax, xv_c + bw, yv4 + 0.08, xht, y2 + bh - 0.10)
    _arrow(ax, xv + bw, y2 + 0.95 + bh / 2, xht, y2 + bh / 2)
    _arrow(ax, xv + bw, y2 - 0.95 + bh / 2, xht, y2 + 0.10)
    _arrow(ax, xv + bw, y2 + 0.95 + 0.10, xht, y3 + bh - 0.08)
    _arrow(ax, xv + bw, y3 + bh / 2, xht, y3 + bh / 2)

    _rect(ax, xout, 2.15, 1.78, 3.55, "")
    ax.text(xout + 0.89, yc, "ВЫВОД\nH1, H2, H3 (t(k))\nV1–V7\nP6–P11", ha="center", va="center", fontsize=8)
    _arrow(ax, xht + 1.42, y1 + bh / 2, xout, 5.15)
    _arrow(ax, xht + 1.42, y2 + bh / 2, xout, yc)
    _arrow(ax, xht + 1.42, y3 + bh / 2, xout, 2.70)

    ax.add_patch(mp.Ellipse((16.15, yc), 1.20, 0.82, fill=False, lw=1.2))
    ax.text(16.15, yc, "Стоп", ha="center", va="center", fontsize=10)
    _arrow(ax, xout + 1.78, yc, 15.55, yc)

    fig.tight_layout()
    fig.savefig(path, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_doc(scheme: Path, matrix: Path, flow: Path) -> None:
    doc = Document()
    doc.element.set("{http://www.w3.org/2000/xmlns/}m", M_NS)
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for t in [
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное бюджетное образовательное учреждение высшего образования",
        "«Российский химико-технологический университет им. Д.И. Менделеева»",
        "Кафедра информатики и компьютерного проектирования",
    ]:
        add_p(doc, t, size=12, align="center", space_after=0)

    add_p(doc, "", size=12, space_after=16)
    add_p(doc, "ОТЧЕТ", size=16, bold=True, align="center", space_after=8)
    add_p(doc, "По лабораторной работе №1 на тему:", size=14, bold=True, align="center", space_after=4)
    add_p(
        doc,
        "«Моделирование простой гидравлической системы в статическом режиме»",
        size=14,
        align="center",
        space_after=8,
    )
    add_p(doc, "Вариант – 1", size=14, bold=True, align="center", space_after=16)
    add_p(doc, "Выполнил: ______________________________", size=14, align="right", space_after=2)
    add_p(doc, "Группа: ________________________________", size=14, align="right", space_after=2)
    add_p(doc, "Проверила: _____________________________", size=14, align="right", space_after=16)
    add_p(doc, "г. Москва, 2026", size=14, align="center", space_after=16)

    add_p(doc, "Схематическое изображение гидравлической модели с закрытыми ёмкостями:", size=14, bold=True)
    doc.add_picture(str(scheme), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_p(doc, "Условные обозначения:", size=14, bold=True, space_before=10)
    for line in [
        "P1, P2, P3 – давление жидкости на входе в систему;",
        "P4 – давление жидкости на выходе нижней ветви;",
        "P5 – давление жидкости на выходе из ёмкости 3;",
        "P6, P7, P8 – давление жидкости на дно ёмкости;",
        "P9, P10, P11 – давление газа внутри ёмкости;",
        "V1 – V7 – скорость истечения жидкости через клапан;",
        "k1 – k7 – коэффициент пропускной способности клапана;",
        "H1, H2, H3 – высоты накапливаемой жидкости в ёмкости;",
        "H1G, H2G, H3G – геометрическая высота ёмкости.",
    ]:
        add_p(doc, line, space_after=2)

    add_p(doc, "Допущения", size=14, bold=True, space_before=12)
    for line in [
        "1. Газ идеален.",
        "2. Во всех трубах протекает однофазный поток жидкости, температура которого одинакова на всех участках.",
        "3. Форма закрытой ёмкости цилиндрическая с площадью поперечного сечения S.",
        "4. Все трубы располагаются на одном уровне, в системе нет рециклических потоков, или рециклов, не учитываются местные сопротивления и перепады давлений в трубах, т.е. рассматриваются, так называемые короткие трубопроводы.",
    ]:
        add_p(doc, line, align="justify", space_after=3)
    p5 = add_p(doc, "5. В емкостях, не занятых жидкостью, давление газа ", align="justify", space_after=3)
    append_runs(p5, [omml_pn(), " * H1G / (H1G – H1)."])
    add_p(
        doc,
        "6. Системы включают только клапаны(вентили) с постоянными, не изменяющимися коэффициентами пропускной способности и закрытые емкости.",
        align="justify",
        space_after=3,
    )

    add_p(doc, "Построение системы уравнений математического описания гидравлической системы:", size=14, bold=True, space_before=12)
    pb = add_p(doc, "ур. Бернулли: V = k * sign (Рвх – Рвых) * ", align="center", italic=True)
    pb.runs[0].italic = True
    append_runs(pb, [omml_sqrt_abs("вх", "вых")])
    add_p(doc, "sign – функция знака: -1, 0, +1", align="center", space_after=8)
    add_p(doc, "ур. Паскаля: Рж = Рг + ρ*g* H", align="center", italic=True, space_after=10)

    add_p(doc, "Определение скорости потоков жидкости:", size=14, bold=True)
    for vi, a, b in [
        ("V1 = k1 * sign (Р1 – Р6) * ", "1", "6"),
        ("V2 = k2 * sign (Р2 – Р6) * ", "2", "6"),
        ("V3 = k3 * sign (Р3 – Р6) * ", "3", "6"),
        ("V4 = k4 * sign (Р6 – Р7) * ", "6", "7"),
        ("V5 = k5 * sign (Р7 – Р8) * ", "7", "8"),
        ("V6 = k6 * sign (Р7 – Р4) * ", "7", "4"),
        ("V7 = k7 * sign (Р8 – Р5) * ", "8", "5"),
    ]:
        pv = add_p(doc, vi, space_after=3)
        append_runs(pv, [omml_sqrt_abs(a, b)])

    add_p(doc, "Уравнения массового баланса:", size=14, bold=True, space_before=10)
    add_p(doc, "8’) (V1 + V2 + V3 – V4 )/S1= (H1(t(k))- H1(t(o)))/ Δt = f8", space_after=2)
    add_p(doc, "8*) H1(t(o))= H1(o)", space_after=6)
    add_p(doc, "9’) (V4 – V5 – V6 )/S2= (H2(t(k))- H2(t(o)))/ Δt = f9", space_after=2)
    add_p(doc, "9*) H2(t(o))= H2(o)", space_after=6)
    add_p(doc, "10’) (V5 – V7 )/S3= (H3(t(k))- H3(t(o)))/ Δt = f10", space_after=2)
    add_p(doc, "10*) H3(t(o))= H3(o)", space_after=8)

    add_p(doc, "Определение давлений жидкости и газа в закрытых ёмкостях:", size=14, bold=True, space_before=10)
    for line in [
        "11) Р6 = Р9 + ρ*g* H1",
        "12) Р9 = РN * H1G/ (H1G – H1)",
        "13) Р7 = Р10 + ρ*g* H2",
        "14) Р10 = РN * H2G/ (H2G – H2)",
        "15) Р8 = Р11 + ρ*g* H3",
        "16) Р11 = РN * H3G/ (H3G – H3)",
    ]:
        add_p(doc, line, space_after=3)

    add_p(doc, "Информационная матрица системы уравнений:", size=14, bold=True, space_before=12)
    add_p(doc, "Условные обозначения:", space_after=4)
    add_p(doc, "◇ – определяемая переменная;", space_after=2)
    add_p(doc, "○ – известная переменная (входит в уравнение);", space_after=2)
    add_p(doc, "В столбце «№ шага» отражается последовательность вычислений согласно выбранному алгоритму расчетов.", align="justify")
    doc.add_picture(str(matrix), width=Cm(16.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_p(doc, "Блок-схема алгоритма расчётов стационарного режима гидравлической системы:", size=14, bold=True, space_before=12)
    doc.add_picture(str(flow), width=Cm(16.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_p(
        doc,
        "Компьютерная программа на языке MATLAB в данную лабораторную работу не входит.",
        italic=True,
        space_before=12,
    )

    doc.save(DOCX)


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    scheme = OUT_DIR / "схема_вариант1.png"
    shutil.copy2(USER_SCHEME, scheme)
    matrix = OUT_DIR / "информационная_матрица.png"
    flow = OUT_DIR / "блок_схема.png"
    draw_matrix(matrix)
    draw_flow(flow)
    build_doc(scheme, matrix, flow)
    print("saved", DOCX.resolve())


if __name__ == "__main__":
    main()
