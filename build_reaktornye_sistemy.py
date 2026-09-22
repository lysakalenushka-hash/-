#!/usr/bin/env python3
"""Л/р «Реакторные системы». Вариант 2. n1 = n2 = 1. Четыре схемы по 4 реактора."""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("/workspace/reaktory_sistemy_variant2/Реакторные_системы_вариант2.xlsx")

CA0 = 30.0
V_SYS = 100.0  # л, суммарный объём системы
V0_SYS = 50.0  # л/мин через систему
N = 4
V_I = V_SYS / N  # 25 л
K1, K2 = 0.6, 0.4  # мин⁻¹, как л/р №3 случай n1 = n2 = 1
PI_FACTOR = V0_SYS * 60 / 1000.0  # ΠR = 3·CR кмоль/ч при полном расходе системы

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
SUB_FONT = Font(name="Calibri", bold=True, size=12, color="2E75B6")
CELL_FONT = Font(name="Calibri", size=11)
NUM_FMT = "0.0000"
CONC_FMT = "0.00"


def thin():
    return Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )


def style_header(cell):
    cell.fill = HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin()


def style_cell(cell, fmt=None):
    cell.font = CELL_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin()
    if fmt:
        cell.number_format = fmt


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def xyz_from_c(ca, cr, cs):
    x = 1.0 - ca
    y = cr
    z = cs
    s = (1.0 - K2 * cr / (K1 * ca)) if ca > 1e-12 else 0.0
    return x, y, z, s


def rec(tau, x, y, z):
    return {
        "tau": tau,
        "X": x,
        "CA": CA0 * (1.0 - x),
        "Y": y,
        "CR": CA0 * y,
        "Z": z,
        "CS": CA0 * z,
        "PiR": PI_FACTOR * CA0 * y,
    }


# --- РИС-н в ряду: четыре экрана (τ_i = 0,5 мин) ---
CSTR_SERIES = {
    1: [
        (0.0, 0.0000, 0.0000, 0.0000),
        (0.5, 0.2308, 0.1923, 0.0385),
    ],
    2: [
        (0.0, 0.2308, 0.1923, 0.0385),
        (0.5, 0.4083, 0.3082, 0.1001),
    ],
    3: [
        (0.0, 0.4083, 0.3082, 0.1001),
        (0.5, 0.5448, 0.3706, 0.1742),
    ],
    4: [
        (0.0, 0.5448, 0.3706, 0.1742),
        (0.5, 0.6498, 0.3964, 0.2535),
    ],
}

# л/р №3, единичные, n1 = n2 = 1, τ = 2 мин
SINGLE_CSTR = (2.0, 0.5454, 0.3030, 0.2424)
SINGLE_PFR = (2.0, 0.6988, 0.4444, 0.2544)
# выходы четырёх РИВ в ряду (τi = 0,5 мин), X Y Z
PFR_SERIES_OUT = [
    (0.2592, 0.2337, 0.0254),
    (0.4512, 0.3645, 0.0867),
    (0.5934, 0.4267, 0.1667),
    (0.6988, 0.4444, 0.2544),
]


def cstr_from_inlet(tau, ca_in, cr_in, cs_in):
    """Один РИС-н, безразмерные концентрации (C0 = 1)."""
    ca = ca_in / (1.0 + K1 * tau)
    cr = (cr_in + K1 * ca * tau) / (1.0 + K2 * tau)
    cs = cs_in + K2 * cr * tau
    return ca, cr, cs


def pfr_step(tau, ca_in, cr_in, cs_in):
    if tau == 0:
        return ca_in, cr_in, cs_in
    ca = ca_in * math.exp(-K1 * tau)
    cr = cr_in * math.exp(-K2 * tau) + (K1 / (K2 - K1)) * ca_in * (
        math.exp(-K1 * tau) - math.exp(-K2 * tau)
    )
    cs = ca_in + cr_in + cs_in - ca - cr
    return ca, cr, cs


def pfr_profile(tau_max, n_steps, ca_in=1.0, cr_in=0.0, cs_in=0.0):
    rows = []
    for i in range(n_steps + 1):
        t = round(tau_max * i / n_steps, 4)
        ca, cr, cs = pfr_step(t, ca_in, cr_in, cs_in)
        x, y, z, _s = xyz_from_c(ca, cr, cs)
        rows.append((t, round(x, 4), round(y, 4), round(z, 4)))
    return rows, ca, cr, cs


SCHEME_FILL = PatternFill("solid", fgColor="2E75B6")
SCHEME_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
PUNKT_FONT = Font(name="Calibri", bold=True, size=13, color="1F4E79")


def write_punkt(ws, r0, text):
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=7)
    cell = ws.cell(r0, 1, text)
    cell.font = PUNKT_FONT
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 22
    return r0 + 1


def write_scheme(ws, r0, text):
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=7)
    cell = ws.cell(r0, 1, text)
    cell.font = SCHEME_FONT
    cell.fill = SCHEME_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 24
    return r0 + 1


def write_subst(ws, r0, title, x, y, z):
    """Три строки CA, CR, CS с подстановкой, как просит отчёт."""
    r0 = write_punkt(ws, r0, title)
    for i, h in enumerate(
        ["Величина", "Формула", "С экрана", "Подстановка", "Результат, моль/л"],
        1,
    ):
        style_header(ws.cell(r0, i, h))
    ws.row_dimensions[r0].height = 24
    lines = [
        ("CA, моль/л", "CA0·(1−X)", f"X = {x:.4f}".replace(".", ","),
         f"30·(1−{x:.4f})".replace(".", ","), CA0 * (1.0 - x)),
        ("CR, моль/л", "CA0·Y", f"Y = {y:.4f}".replace(".", ","),
         f"30·{y:.4f}".replace(".", ","), CA0 * y),
        ("CS, моль/л", "CA0·Z", f"Z = {z:.4f}".replace(".", ","),
         f"30·{z:.4f}".replace(".", ","), CA0 * z),
    ]
    for i, (a, b, c, d, e) in enumerate(lines):
        style_cell(ws.cell(r0 + 1 + i, 1, a))
        style_cell(ws.cell(r0 + 1 + i, 2, b))
        style_cell(ws.cell(r0 + 1 + i, 3, c))
        style_cell(ws.cell(r0 + 1 + i, 4, d))
        cell = ws.cell(r0 + 1 + i, 5, e)
        style_cell(cell, CONC_FMT)
    return r0 + 5


def write_block(ws, r0, title, rows):
    headers = [
        "τ, мин",
        "X",
        "CA, моль/л",
        "Y",
        "CR, моль/л",
        "Z",
        "CS, моль/л",
    ]
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=7)
    t = ws.cell(r0, 1, title)
    t.font = SUB_FONT
    t.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 22
    hr = r0 + 1
    for i, h in enumerate(headers):
        style_header(ws.cell(hr, i + 1, h))
    ws.row_dimensions[hr].height = 28
    for i, raw in enumerate(rows):
        d = rec(*raw)
        vals = [d["tau"], d["X"], d["CA"], d["Y"], d["CR"], d["Z"], d["CS"]]
        fmts = ["0.00", NUM_FMT, CONC_FMT, NUM_FMT, CONC_FMT, NUM_FMT, CONC_FMT]
        for j, (v, fmt) in enumerate(zip(vals, fmts), 1):
            cell = ws.cell(hr + 1 + i, j, v)
            style_cell(cell, fmt)
        if i % 2 == 1:
            for j in range(1, 8):
                ws.cell(hr + 1 + i, j).fill = PatternFill("solid", fgColor="D6EAF8")
    return hr + 1 + len(rows)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    widths7 = [14, 12, 16, 12, 16, 12, 16]

    ca_p, cr_p, cs_p = cstr_from_inlet(2.0, 1.0, 0.0, 0.0)
    x_par, y_par, z_par, _ = xyz_from_c(ca_p, cr_p, cs_p)
    rows_cstr_par = [
        (0.0, 0.0000, 0.0000, 0.0000),
        (2.0, 0.5454, 0.3030, 0.2424),
    ]
    ca_in, cr_in, cs_in = 1.0, 0.0, 0.0
    pfr_series = []
    for _i in range(4):
        rows, ca_in, cr_in, cs_in = pfr_profile(0.5, 5, ca_in, cr_in, cs_in)
        pfr_series.append(rows)
    rows_pfr_par, *_rest = pfr_profile(2.0, 10, 1.0, 0.0, 0.0)

    # ========== п. 1 ==========
    w1 = wb.active
    w1.title = "1. Концентрации"
    w1["A1"] = "Пункт 1. Расчёт реальных концентраций веществ A, R, S. Вариант 2"
    w1["A1"].font = TITLE_FONT
    w1.merge_cells("A1:H1")
    w1.row_dimensions[1].height = 24

    write_punkt(
        w1,
        3,
        "Программа выдаёт безразмерные X, Y, Z (отнесённые к CA0). Реальные концентрации:",
    )
    w1.merge_cells("A4:H8")
    w1["A4"] = (
        "CA = CA0 · (1 − X)     — непрореагировавший реагент A\n"
        "CR = CA0 · Y          — целевой продукт R\n"
        "CS = CA0 · Z          — побочный продукт S\n"
        "CA0 = 30 моль/л. На экране 8 строк (вход и выход каждого из 4 РИС-н), "
        "но вход следующего аппарата совпадает с выходом предыдущего, поэтому уникальных состояний пять."
    )
    w1["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    w1["A4"].font = Font(name="Calibri", size=12)
    for rr in (4, 5, 6, 7, 8):
        w1.row_dimensions[rr].height = 18

    subst_blocks = [
        ("Вход в систему / вход реактора 1 (τ = 0):", 0.0000, 0.0000, 0.0000),
        ("Выход реактора 1 (τ = 0,5) — он же вход реактора 2:", 0.2308, 0.1923, 0.0385),
        ("Выход реактора 2 (τ = 0,5) — он же вход реактора 3:", 0.4083, 0.3082, 0.1001),
        ("Выход реактора 3 (τ = 0,5) — он же вход реактора 4:", 0.5448, 0.3706, 0.1742),
        ("Выход реактора 4 (τ = 0,5) — выход системы 1:", 0.6498, 0.3964, 0.2535),
    ]
    r = 9
    for title, x, y, z in subst_blocks:
        r = write_subst(w1, r, title, x, y, z) + 1

    w1.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=8)
    w1.cell(
        r,
        1,
        "Проверка: CA + CR + CS = CA0, то есть (1−X) + Y + Z = 1. "
        "Расхождения в 0,01 моль/л — округление X, Y, Z на экране. "
        "Дальше все таблицы пункта 2 считаются этими же формулами.",
    ).alignment = Alignment(wrap_text=True, vertical="top")
    w1.cell(r, 1).font = Font(name="Calibri", size=12)

    set_widths(w1, [48, 18, 18, 22, 22, 16, 16, 16])
    w1.page_setup.orientation = "landscape"
    w1.page_setup.fitToPage = True
    w1.page_setup.fitToWidth = 1
    w1.page_setup.fitToHeight = 0

    # ========== п. 2 ==========
    w2 = wb.create_sheet("2. Таблицы")
    w2["A1"] = "Пункт 2. Таблицы моделирования для каждого реактора каждой системы"
    w2["A1"].font = TITLE_FONT
    w2.merge_cells("A1:G1")
    r = 3

    r = write_scheme(
        w2,
        r,
        "Система 1. Четыре последовательно соединённых РИС-н одинакового объёма  "
        "(vi = 50 л/мин, τi = 0,5 мин). Экран программы.",
    )
    r += 1
    for i in range(1, 5):
        r = write_block(w2, r, f"Реактор {i}", CSTR_SERIES[i]) + 2

    r = write_scheme(
        w2,
        r,
        "Система 2. Четыре параллельно соединённых РИС-н одинакового объёма, "
        "одинаковая нагрузка (vi = 12,5 л/мин, τi = 2 мин).",
    )
    r += 1
    for i in range(1, 5):
        r = write_block(w2, r, f"Реактор {i}", rows_cstr_par) + 2

    r = write_scheme(
        w2,
        r,
        "Система 3. Четыре последовательно соединённых РИВ одинакового объёма "
        "(vi = 50 л/мин, τi = 0,5 мин).",
    )
    r += 1
    for i, rows in enumerate(pfr_series, 1):
        r = write_block(w2, r, f"Реактор {i}", rows) + 2

    r = write_scheme(
        w2,
        r,
        "Система 4. Четыре параллельно соединённых РИВ одинакового объёма, "
        "одинаковая нагрузка (vi = 12,5 л/мин, τi = 2 мин).",
    )
    r += 1
    for i in range(1, 5):
        r = write_block(w2, r, f"Реактор {i}", rows_pfr_par) + 2

    set_widths(w2, widths7)
    w2.page_setup.orientation = "landscape"
    w2.page_setup.fitToPage = True
    w2.page_setup.fitToWidth = 1
    w2.page_setup.fitToHeight = 0
    w2.sheet_properties.pageSetUpPr.fitToPage = True
    w2.freeze_panes = "A3"
    w2.print_title_rows = "1:1"

    # ========== п. 3 ==========
    w3 = wb.create_sheet("3. Производительность")
    w3["A1"] = (
        "Пункт 3. Производительность систем по целевому продукту R "
        "и сравнение с единичными РИС-н и РИВ (л/р №3, n1 = n2)"
    )
    w3["A1"].font = TITLE_FONT
    w3.merge_cells("A1:F1")
    w3.row_dimensions[1].height = 32
    w3.merge_cells("A3:F3")
    w3["A3"] = (
        "ΠR = v0 · CR_вых. v0 системы = 50 л/мин = 3 м³/ч, CR в кмоль/м³ (= моль/л), "
        "поэтому ΠR = 3 · CR, кмоль/ч. Для параллельных схем смешение потоков: "
        "CR_вых = CR каждого аппарата (нагрузки равны)."
    )
    w3["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    w3.row_dimensions[3].height = 36

    headers = ["Схема", "X вых.", "Y вых.", "CR, моль/л", "ΠR, кмоль/ч", "Сравнение с л/р №3"]
    for i, h in enumerate(headers, 1):
        style_header(w3.cell(5, i, h))
    w3.row_dimensions[5].height = 32

    y_ser = CSTR_SERIES[4][-1][2]
    y_cstr = SINGLE_CSTR[2]
    y_pfr = SINGLE_PFR[2]
    data = [
        ("Единичный РИС-н, л/р №3", SINGLE_CSTR[1], y_cstr, "база для схем 2"),
        ("Единичный РИВ, л/р №3", SINGLE_PFR[1], y_pfr, "база для схем 3 и 4"),
        ("1. 4 РИС-н последовательно", CSTR_SERIES[4][-1][1], y_ser, "выше единичного РИС-н, ниже РИВ"),
        ("2. 4 РИС-н параллельно", SINGLE_CSTR[1], y_cstr, "совпадает с единичным РИС-н"),
        ("3. 4 РИВ последовательно", SINGLE_PFR[1], y_pfr, "совпадает с единичным РИВ"),
        ("4. 4 РИВ параллельно", SINGLE_PFR[1], y_pfr, "совпадает с единичным РИВ"),
    ]
    for i, (name, x, y, note) in enumerate(data):
        cr = CA0 * y
        pir = PI_FACTOR * cr
        vals = [name, x, y, cr, pir, note]
        fmts = [None, NUM_FMT, NUM_FMT, CONC_FMT, "0.00", None]
        for j, (v, fmt) in enumerate(zip(vals, fmts), 1):
            cell = w3.cell(6 + i, j, v)
            style_cell(cell, fmt)
            if j in (1, 6):
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if i >= 2:
            for j in range(1, 7):
                w3.cell(6 + i, j).fill = PatternFill("solid", fgColor="FFF2CC")
        if "РИВ" in name:
            for j in range(1, 7):
                w3.cell(6 + i, j).fill = PatternFill("solid", fgColor="E2EFDA")

    w3.merge_cells("A13:F17")
    w3["A13"] = (
        "Почему так. Параллельные РИС-н: каждый получает v0/4 и Vi = V/4, поэтому τi = Vi/vi = 2 мин — "
        "как у единичного РИС-н, выход и ΠR те же (27,27 кмоль/ч). "
        "Параллельные РИВ: τi = 2 мин, выход как у единичного РИВ, ΠR = 40,00 кмоль/ч. "
        "Последовательные РИВ: времена пребывания складываются, Στ = 2 мин — снова один РИВ, ΠR = 40,00. "
        "Последовательные РИС-н: каскад четырёх ёмкостей приближает режим к вытеснению, "
        "ΠR = 35,68 кмоль/ч — между 27,27 и 40,00."
    )
    w3["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    w3["A13"].font = Font(name="Calibri", size=12)
    w3.row_dimensions[13].height = 48

    chart = BarChart()
    chart.type = "col"
    chart.title = "ΠR, кмоль/ч"
    chart.y_axis.title = "ΠR, кмоль/ч"
    chart.legend = None
    data_ref = Reference(w3, min_col=5, min_row=5, max_row=11)
    cats = Reference(w3, min_col=1, min_row=6, max_row=11)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 18
    chart.height = 9
    w3.add_chart(chart, "A19")
    set_widths(w3, [34, 12, 12, 14, 16, 42])
    w3.page_setup.orientation = "landscape"

    # ========== п. 4 ==========
    w4 = wb.create_sheet("4. Сравнение систем")
    w4["A1"] = (
        "Пункт 4. Сравнение реакторных систем по производительности и по гидравлическому сопротивлению"
    )
    w4["A1"].font = TITLE_FONT
    w4.merge_cells("A1:D1")
    w4.row_dimensions[1].height = 28

    h4 = ["Система", "ΠR, кмоль/ч", "Гидравлическое сопротивление", "Итог"]
    for i, h in enumerate(h4, 1):
        style_header(w4.cell(3, i, h))
    w4.row_dimensions[3].height = 32
    rows4 = [
        (
            "1. 4 РИС-н последовательно",
            35.68,
            "Через каждый аппарат идёт полный расход 50 л/мин, четыре перепада складываются — ΔP высокое.",
            "ΠR средняя, гидравлика тяжёлая",
        ),
        (
            "2. 4 РИС-н параллельно",
            27.27,
            "Расход на ветвь 12,5 л/мин, общий ΔP ≈ одной ветви — сопротивление низкое.",
            "ΠR как у одного РИС-н, гидравлика лёгкая",
        ),
        (
            "3. 4 РИВ последовательно",
            40.00,
            "Полный расход 50 л/мин через все четыре РИВ, ΔP максимальное среди схем.",
            "ΠR максимальная, гидравлика самая тяжёлая",
        ),
        (
            "4. 4 РИВ параллельно",
            40.00,
            "Расход 12,5 л/мин на ветвь, ΔP как у одного РИВ малого расхода — ниже, чем в ряду.",
            "ΠR максимальная при меньшем ΔP — предпочтительная схема",
        ),
    ]
    for i, (a, b, c, d) in enumerate(rows4):
        style_cell(w4.cell(4 + i, 1, a))
        w4.cell(4 + i, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cellb = w4.cell(4 + i, 2, b)
        style_cell(cellb, "0.00")
        style_cell(w4.cell(4 + i, 3, c))
        w4.cell(4 + i, 3).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        style_cell(w4.cell(4 + i, 4, d))
        w4.cell(4 + i, 4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        w4.row_dimensions[4 + i].height = 48
        if i in (2, 3):
            for j in range(1, 5):
                w4.cell(4 + i, j).fill = PatternFill("solid", fgColor="E2EFDA")

    w4.merge_cells("A9:D12")
    w4["A9"] = (
        "Вывод по п. 4. По производительности системы ранжируются так: "
        "4 РИВ (ряд или параллель, 40,00) > 4 РИС-н в ряду (35,68) > 4 РИС-н параллельно (27,27). "
        "По гидравлическому сопротивлению наоборот выгоднее параллель: меньший расход на аппарат. "
        "Компромисс — четыре РИВ параллельно: та же ΠR, что у вытеснения, при меньшем сопротивлении, "
        "чем у четырёх РИВ в ряду."
    )
    w4["A9"].alignment = Alignment(wrap_text=True, vertical="top")
    w4["A9"].font = Font(name="Calibri", size=12)
    w4.row_dimensions[9].height = 40
    set_widths(w4, [34, 16, 70, 42])
    w4.page_setup.orientation = "landscape"
    w4.page_setup.fitToPage = True
    w4.page_setup.fitToWidth = 1
    w4.page_setup.fitToHeight = 1

    wb.save(OUT)
    print("saved", OUT)
    print("sheets", wb.sheetnames)
    print("PFR series last", pfr_series[-1][-1])


if __name__ == "__main__":
    build()
