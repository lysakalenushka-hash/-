#!/usr/bin/env python3
"""Unify fonts, chrome, footers on v5. Do not touch designer comment blocks."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC = "/home/ubuntu/.cursor/projects/workspace/uploads/Lab-Intellekt_prodayushchaya-prezentaciya_v5_cd47.pptx"
OUT = "/workspace/issledovaniya/Lab-Intellekt_prodayushchaya-prezentaciya_v5.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1A, 0x9B, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF5, 0xF7, 0xFA)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x78)
FOOT = RGBColor(0x9A, 0xA3, 0xB0)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
RED = RGBColor(0xC0, 0x45, 0x3C)
LIGHT = RGBColor(0xC5, 0xD0, 0xDC)
FONT = "Calibri"

ACCENT_W = Emu(109728)
TOP_H = Emu(73152)
KICKER_L, KICKER_T, KICKER_W, KICKER_H = Emu(502920), Emu(201168), Emu(10972800), Emu(292608)
TITLE_L, TITLE_T, TITLE_W, TITLE_H = Emu(502920), Emu(457200), Emu(11155680), Emu(780000)
FOOT_T, FOOT_H = Emu(6565392), Emu(256032)
BRAND_L, BRAND_W = Emu(457200), Emu(8229600)
PAGE_L, PAGE_W = Emu(10424160), Emu(1280160)
BRAND = "Lab-Интеллект  ·  СМАРТА  ·  конфиденциально"

COMMENT_SNIPPETS = (
    "Сюда надо вставить",
    "Сюда вставляем ролик",
    "Этот слайд пока не делаем",
    "Использовать лого исследований",
    "Пришлю отдельно",
)


def is_comment(sh):
    if not sh.has_text_frame:
        return False
    t = sh.text_frame.text
    return any(s in t for s in COMMENT_SNIPPETS)


def set_run(run, size, bold, color, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def restyle_textbox(sh, size, bold, color, italic=False, align=None, anchor=None):
    if not sh.has_text_frame:
        return
    tf = sh.text_frame
    tf.word_wrap = True
    if anchor:
        try:
            tf._txBody.bodyPr.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
        except Exception:
            pass
    for p in tf.paragraphs:
        if align is not None:
            p.alignment = align
        for r in p.runs:
            set_run(r, size, bold, color, italic)


def solid_bg(sh, color):
    try:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
    except Exception:
        pass


def restyle_table(table):
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            header = r == 0
            tf = cell.text_frame
            tf.word_wrap = True
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    set_run(run, 13 if header else 12, header, WHITE if header else INK)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def ensure_footer(slide, page, total, dark=False):
    c = RGBColor(0x8A, 0xA0, 0xB8) if dark else FOOT
    page_box = None
    for sh in slide.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if " /  " in t and len(t) < 18:
            page_box = sh

    def place(box, l, t, w, h, text, align):
        box.left, box.top, box.width, box.height = l, t, w, h
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        set_run(r, 11, False, c)

    if page_box:
        place(page_box, PAGE_L, FOOT_T, PAGE_W, FOOT_H, f"{page}  /  {total}", PP_ALIGN.RIGHT)
    else:
        box = slide.shapes.add_textbox(PAGE_L, FOOT_T, PAGE_W, FOOT_H)
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{page}  /  {total}"
        set_run(r, 11, False, c)


def is_dark_hero(slide):
    for sh in slide.shapes:
        if sh.name == "Rectangle 1":
            xml = etree.tostring(sh._element, encoding="unicode")
            return 'srgbClr val="0B1F3A"' in xml
    return False


def chrome_content(slide, W, H):
    for sh in slide.shapes:
        if is_comment(sh) or sh.has_table:
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            continue
        if sh.name == "Rectangle 1" and sh.width > Emu(10000000) and sh.height > Emu(5000000):
            sh.left, sh.top, sh.width, sh.height = 0, 0, W, H
            solid_bg(sh, OFF)
        elif sh.name == "Rectangle 2" and sh.width < Emu(220000):
            sh.left, sh.top, sh.width, sh.height = 0, 0, ACCENT_W, H
            solid_bg(sh, TEAL)
        elif sh.name == "Rectangle 3" and sh.height < Emu(120000) and sh.top < Emu(30000):
            sh.left, sh.top, sh.width, sh.height = 0, 0, W, TOP_H
            solid_bg(sh, TEAL)


def style_headers(slide):
    for sh in slide.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if abs(sh.top - KICKER_T) < Emu(40000) and sh.left < Emu(700000):
            sh.left, sh.top, sh.width, sh.height = KICKER_L, KICKER_T, KICKER_W, KICKER_H
            if not t:
                p = sh.text_frame.paragraphs[0]
                if p.runs:
                    p.runs[0].text = "КОНТЕКСТ"
                else:
                    r = p.add_run()
                    r.text = "КОНТЕКСТ"
            restyle_textbox(sh, 12, True, TEAL, align=PP_ALIGN.LEFT, anchor="t")
        elif abs(sh.top - TITLE_T) < Emu(80000) and sh.width > Emu(7000000):
            sh.left, sh.top = TITLE_L, TITLE_T
            sh.width = TITLE_W
            if sh.height < Emu(500000):
                sh.height = TITLE_H
            size = 26 if len(t) > 72 else 28
            restyle_textbox(sh, size, True, INK, align=PP_ALIGN.LEFT, anchor="t")


def style_body(slide):
    for sh in slide.shapes:
        if is_comment(sh):
            continue
        if sh.has_table:
            restyle_table(sh.table)
            continue
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if not t:
            continue
        if sh.top >= FOOT_T - Emu(30000):
            continue
        if abs(sh.top - KICKER_T) < Emu(40000) or abs(sh.top - TITLE_T) < Emu(80000):
            continue
        if (len(t) <= 5 and t.endswith("%")) or t in ("20%", "−63%", "10%", "40%"):
            restyle_textbox(sh, 32, True, TEAL, align=PP_ALIGN.LEFT, anchor="ctr")
            continue
        if t in ("01", "02", "03", "1", "2", "3") and sh.height < Emu(600000):
            restyle_textbox(sh, 26, True, TEAL, align=PP_ALIGN.LEFT, anchor="t")
            continue
        if t.startswith("Ваша цифра") or "млн" in t and "₽" in t or "тыс. ₽" in t:
            restyle_textbox(sh, 15, True, RED, align=PP_ALIGN.LEFT, anchor="t")
            continue
        if t.startswith("•"):
            restyle_textbox(sh, 15, False, INK, align=PP_ALIGN.LEFT, anchor="t")
            continue
        if t.startswith(("Gallup", "Метаанализ", "Оценка программы", "Точную цифру")):
            restyle_textbox(sh, 12, False, MUTED, align=PP_ALIGN.LEFT, anchor="t")
            continue
        if sh.height <= Emu(520000) and len(t) < 45 and "\n" not in t:
            restyle_textbox(sh, 18, True, NAVY, align=PP_ALIGN.LEFT, anchor="t")
            continue
        restyle_textbox(sh, 15, False, MUTED, align=PP_ALIGN.LEFT, anchor="t")


def fix_product_grid(slide, W, H):
    """Slide 7: 2x3 feature cards. Comment block untouched."""
    for sh in slide.shapes:
        if is_comment(sh):
            continue
        if sh.name == "Rectangle 1":
            sh.left, sh.top, sh.width, sh.height = 0, 0, W, H
            solid_bg(sh, OFF)

    cols_x = [Emu(502920), Emu(4297680), Emu(8092440)]
    card_w, card_h = Emu(3611880), Emu(1874519)
    row_y = [Emu(1780000), Emu(3780000)]
    pad = Emu(220000)

    cards = [sh for sh in slide.shapes if "Rounded" in sh.name and not is_comment(sh)]
    cards.sort(key=lambda sh: (round(sh.top / 200000), sh.left))
    for i, sh in enumerate(cards[:6]):
        r, c = divmod(i, 3)
        sh.left, sh.top, sh.width, sh.height = cols_x[c], row_y[r], card_w, card_h

    titles = {}
    bodies = {}
    for sh in slide.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        mapping = {
            "Микрокурсы": "t",
            "Программы и назначения": "t",
            "Баллы и вовлеченность": "t",
            "Кабинет компании": "t",
            "Рейтинг": "t",
            "Отчёты": "t",
        }
        if t in mapping:
            titles[t] = sh
        elif t.startswith("Видео, презентации"):
            bodies["Микрокурсы"] = sh
        elif t.startswith("Группы, каждый"):
            bodies["Программы и назначения"] = sh
        elif "методик" in t:
            bodies["Баллы и вовлеченность"] = sh
        elif t.startswith("Штатка"):
            bodies["Кабинет компании"] = sh
        elif "мотиватор" in t:
            bodies["Рейтинг"] = sh
        elif t.startswith("Прогресс, просрочки"):
            bodies["Отчёты"] = sh

    order = [
        ("Микрокурсы", 0, 0),
        ("Программы и назначения", 0, 1),
        ("Баллы и вовлеченность", 0, 2),
        ("Кабинет компании", 1, 0),
        ("Рейтинг", 1, 1),
        ("Отчёты", 1, 2),
    ]
    for name, r, c in order:
        x, y = cols_x[c], row_y[r]
        tb = titles.get(name)
        if tb:
            tb.left = x + pad
            tb.top = y + Emu(200000)
            tb.width = card_w - pad * 2
            tb.height = Emu(420000)
            restyle_textbox(tb, 16, True, NAVY, align=PP_ALIGN.LEFT, anchor="t")
        bd = bodies.get(name)
        if bd:
            bd.left = x + pad
            bd.top = y + Emu(680000)
            bd.width = card_w - pad * 2
            bd.height = Emu(1050000)
            restyle_textbox(bd, 14, False, MUTED, align=PP_ALIGN.LEFT, anchor="t")


def overrides(prs):
    s3 = prs.slides[2]
    for sh in s3.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == "ДВА ФОРМАТА":
            restyle_textbox(sh, 12, True, TEAL, align=PP_ALIGN.LEFT, anchor="t")
        elif t.startswith("Длинный курс"):
            restyle_textbox(sh, 22, True, WHITE, align=PP_ALIGN.LEFT, anchor="t")
        elif t.startswith("5–15"):
            restyle_textbox(sh, 22, True, GOLD, align=PP_ALIGN.LEFT, anchor="t")

    s6 = prs.slides[5]
    for sh in s6.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t in ("Микрообучение", "Аналитика руководителя"):
            restyle_textbox(sh, 18, True, WHITE, align=PP_ALIGN.LEFT, anchor="t")
        if t.startswith("Каждый рабочий") or t.startswith("Контроль просрочки"):
            restyle_textbox(sh, 15, False, LIGHT, align=PP_ALIGN.LEFT, anchor="t")
        if t == "Проверка знания":
            restyle_textbox(sh, 18, True, INK, align=PP_ALIGN.LEFT, anchor="t")

    s8 = prs.slides[7]
    for sh in s8.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == "Моё обучение":
            restyle_textbox(sh, 22, True, WHITE, align=PP_ALIGN.LEFT, anchor="t")
        if t.startswith("Микрообучение"):
            restyle_textbox(sh, 18, False, WHITE, align=PP_ALIGN.LEFT, anchor="t")

    s10 = prs.slides[9]
    for sh in s10.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t in ("СДО", "Очное обучение", "Микрообучение"):
            restyle_textbox(sh, 20, True, WHITE, align=PP_ALIGN.CENTER, anchor="ctr")
        elif t in ("Система дистанционного обучения", "Класс и преподаватель", "Диалог и выбор ответа"):
            restyle_textbox(sh, 12, False, LIGHT, align=PP_ALIGN.CENTER, anchor="t")

    s15 = prs.slides[14]
    for sh in s15.shapes:
        if is_comment(sh) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == "Бизнес":
            restyle_textbox(sh, 20, True, TEAL, align=PP_ALIGN.LEFT, anchor="t")
        if t == "до 500 человек":
            restyle_textbox(sh, 14, True, GOLD, align=PP_ALIGN.LEFT, anchor="t")
        if t.startswith("Безлимит"):
            restyle_textbox(sh, 15, False, LIGHT, align=PP_ALIGN.LEFT, anchor="t")
        if t in ("Старт", "Корпоратив"):
            restyle_textbox(sh, 20, True, NAVY, align=PP_ALIGN.LEFT, anchor="t")


def main():
    prs = Presentation(SRC)
    W, H = prs.slide_width, prs.slide_height
    total = len(prs.slides)

    fix_product_grid(prs.slides[6], W, H)

    for idx, slide in enumerate(prs.slides):
        page = idx + 1
        dark = is_dark_hero(slide)
        if dark:
            for sh in slide.shapes:
                if is_comment(sh):
                    continue
                if sh.name == "Rectangle 1":
                    sh.left, sh.top, sh.width, sh.height = 0, 0, W, H
                if sh.name == "Rectangle 2":
                    sh.left, sh.top, sh.width, sh.height = 0, 0, Emu(164592), H
                    solid_bg(sh, TEAL)
                if sh.name == "Rectangle 3" and sh.top > Emu(5000000):
                    sh.left, sh.top, sh.width, sh.height = 0, Emu(6537960), W, Emu(320040)
                    solid_bg(sh, TEAL)
                if sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if "ДЕМО" in t:
                        restyle_textbox(sh, 14, True, GOLD, align=PP_ALIGN.LEFT, anchor="t")
                    elif t.startswith("Смотрим") or (t.startswith("Микрообучение") and sh.height > Emu(800000)):
                        restyle_textbox(sh, 32, True, WHITE, align=PP_ALIGN.LEFT, anchor="t")
                    elif t.startswith("1."):
                        restyle_textbox(sh, 18, False, LIGHT, align=PP_ALIGN.LEFT, anchor="t")
            ensure_footer(slide, page, total, dark=True)
        else:
            chrome_content(slide, W, H)
            style_headers(slide)
            style_body(slide)
            ensure_footer(slide, page, total, dark=False)

    overrides(prs)
    prs.save(OUT)
    print("saved", OUT)
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if is_comment(sh):
                print("comment intact", i, sh.left, sh.top, sh.width, sh.height)


if __name__ == "__main__":
    main()
