#!/usr/bin/env python3
"""Build v9 from v5 and insert a marketing analytics slide after «Глазами ОТ / HR»."""

import shutil
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC = "/workspace/issledovaniya/Lab-Intellekt_prodayushchaya-prezentaciya_v5.pptx"
OUT = "/workspace/issledovaniya/Lab-Intellekt_prodayushchaya-prezentaciya_v9.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1A, 0x9B, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF5, 0xF7, 0xFA)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x78)
FOOT = RGBColor(0x9A, 0xA3, 0xB0)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
RED = RGBColor(0xC0, 0x45, 0x3C)
FONT = "Calibri"

COMMENT_SNIPPETS = (
    "Сюда надо вставить",
    "Сюда вставляем ролик",
    "Этот слайд пока не делаем",
    "Использовать лого исследований",
    "Пришлю отдельно",
)

REPORTS = [
    {
        "tag": "Компания",
        "title": "План, факт и прогноз",
        "body": "Одна траектория на подразделение: сколько микроуроков должны пройти, сколько прошли и куда придёте при текущем темпе. Отставание видно до срыва срока, а не после проверки знаний.",
        "accent": TEAL,
    },
    {
        "tag": "Люди",
        "title": "Кто учится — кто только кликает",
        "body": "Рейтинг и четыре сигнала по каждому: активность, качество ответов, просрочка, вовлечённость. Куратор идёт не «ко всем», а к тем, кто в системе каждый день, но отвечает наугад.",
        "accent": GOLD,
    },
    {
        "tag": "Даты",
        "title": "Когда учатся на самом деле",
        "body": "Тепловая карта смены: утро, день, выходные. Микроуроки внутри рабочего дня, без отрыва в класс. Если пик ночью или пустые дни — это уже не обучение, а имитация.",
        "accent": NAVY,
    },
    {
        "tag": "Контент",
        "title": "Где навык, где провал",
        "body": "Программа рядом с программой: где 80% правильных, где копится просрочка. Дорабатываете конкретный урок, а не переписываете весь курс «на всякий случай».",
        "accent": RED,
    },
]


def is_comment(sh):
    if not sh.has_text_frame:
        return False
    return any(s in sh.text_frame.text for s in COMMENT_SNIPPETS)


def set_run(run, size=18, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return box


def move_slide(prs, old_index, new_index):
    sld_id_lst = prs.slides._sldIdLst
    el = list(sld_id_lst)[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def retitle_footers(prs, total):
    for idx, slide in enumerate(prs.slides):
        page = idx + 1
        for sh in slide.shapes:
            if is_comment(sh) or not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if " /  " in t and len(t) < 18:
                tf = sh.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                r = p.add_run()
                r.text = f"{page}  /  {total}"
                set_run(r, size=11, color=FOOT)


def build_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    add_rect(s, 0, 0, W, H, OFF)
    add_rect(s, 0, 0, Emu(109728), H, TEAL)
    add_rect(s, 0, 0, W, Emu(73152), TEAL)

    tb(s, Emu(502920), Emu(201168), Emu(10972800), Emu(292608), "АНАЛИТИКА", size=12, bold=True, color=TEAL)
    tb(
        s,
        Emu(502920),
        Emu(430000),
        Emu(11155680),
        Emu(720000),
        "Четыре отчёта. Сразу видно: кто учится, а кто только «проходит»",
        size=26,
        bold=True,
        color=INK,
    )

    cols = [Emu(502920), Emu(6100000)]
    rows = [Emu(1220000), Emu(3380000)]
    cw, ch = Emu(5400000), Emu(1980000)
    pad = Emu(200000)

    for i, rep in enumerate(REPORTS):
        r, c = divmod(i, 2)
        x, y = cols[c], rows[r]
        add_round(s, x, y, cw, ch, WHITE)
        add_rect(s, x, y, Emu(82000), ch, rep["accent"])
        tb(s, x + pad, y + Emu(140000), cw - pad * 2, Emu(280000), rep["tag"].upper(), size=12, bold=True, color=rep["accent"])
        tb(s, x + pad, y + Emu(420000), cw - pad * 2, Emu(400000), rep["title"], size=18, bold=True, color=NAVY)
        tb(s, x + pad, y + Emu(860000), cw - pad * 2, Emu(980000), rep["body"], size=14, color=MUTED)

    add_round(s, Emu(502920), Emu(5480000), Emu(11155680), Emu(960000), WHITE)
    tb(
        s,
        Emu(731520),
        Emu(5600000),
        Emu(10700000),
        Emu(740000),
        "LMS ставит галочку «прошёл». Lab-Интеллект показывает разрыв: высокая активность при слабых ответах и отставании от плана. Это и есть точка управления охраной труда — не после НС, а пока ещё можно скорректировать назначения и контент.",
        size=15,
        color=INK,
    )

    tb(s, Emu(10424160), Emu(6565392), Emu(1280160), Emu(256032), "10  /  18", size=11, color=FOOT, align=PP_ALIGN.RIGHT)
    return s


def main():
    shutil.copy2(SRC, OUT)
    prs = Presentation(OUT)
    already = False
    for s in prs.slides:
        texts = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if "Четыре отчёта. Сразу видно" in texts:
            already = True
            break
    if not already:
        build_slide(prs)
        # new slide is last; place after «Глазами ОТ / HR» (index 8 → new index 9)
        move_slide(prs, len(prs.slides) - 1, 9)
    retitle_footers(prs, len(prs.slides))
    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides))
    for i, s in enumerate(prs.slides, 1):
        t = ""
        for sh in s.shapes:
            if sh.has_text_frame and sh.top < 900000 and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().split("\n")[0][:70]
                break
        print(f"{i:02d}", t)


if __name__ == "__main__":
    main()
