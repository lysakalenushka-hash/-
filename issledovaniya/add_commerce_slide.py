#!/usr/bin/env python3
"""Append the last commerce slide: web prices from the SMARTA proposal, server unpriced."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC = "/workspace/issledovaniya/Lab-Intellekt_prodayushchaya-prezentaciya_v5.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1A, 0x9B, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF5, 0xF7, 0xFA)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x78)
FOOT = RGBColor(0x9A, 0xA3, 0xB0)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
FONT = "Calibri"

COMMENT_SNIPPETS = (
    "Сюда надо вставить",
    "Сюда вставляем ролик",
    "Этот слайд пока не делаем",
    "Использовать лого исследований",
    "Пришлю отдельно",
)

# From «Предложение_технология микрообучения_СМАРТА»: web-version, per worker.
WEB_PRICES = [
    ("Количество пользователей", "Цена за 1 работника"),
    ("0–50", "490 ₽"),
    ("50–300", "460 ₽"),
    ("300–1000", "420 ₽"),
    ("1000–5000", "390 ₽"),
    ("5000 и выше", "договорная"),
]


def is_comment(sh):
    if not sh.has_text_frame:
        return False
    t = sh.text_frame.text
    return any(s in t for s in COMMENT_SNIPPETS)


def set_run(run, size=18, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_round(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return box


def bullets(slide, l, t, w, h, items, size=15, color=INK, spacing=12):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)
    return box


def add_price_table(slide, l, t, w, h, rows):
    shape = slide.shapes.add_table(len(rows), 2, l, t, w, h)
    table = shape.table
    table.columns[0].width = int(w * 0.56)
    table.columns[1].width = int(w * 0.44)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = val
            header = r == 0
            set_run(run, size=13, bold=header or c == 1, color=WHITE if header else INK)
            cell.fill.solid()
            if header:
                cell.fill.fore_color.rgb = TEAL
            elif r % 2 == 0:
                cell.fill.fore_color.rgb = OFF
            else:
                cell.fill.fore_color.rgb = WHITE
    return shape


def retitle_footers(prs, total):
    for idx, slide in enumerate(prs.slides):
        page = idx + 1
        for sh in slide.shapes:
            if is_comment(sh) or not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if " /  " in t and len(t) < 18:
                tf = sh.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                r = p.add_run()
                r.text = f"{page}  /  {total}"
                set_run(r, size=11, color=FOOT)


def build_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    add_rect(s, 0, 0, W, H, OFF)
    add_rect(s, 0, 0, Emu(109728), H, TEAL)
    add_rect(s, 0, 0, W, Emu(73152), TEAL)

    tb(s, Emu(502920), Emu(201168), Emu(10972800), Emu(292608), "КОММЕРЦИЯ", size=12, bold=True, color=TEAL)
    tb(
        s,
        Emu(502920),
        Emu(457200),
        Emu(11155680),
        Emu(720000),
        "Веб-версия и серверная версия",
        size=28,
        bold=True,
        color=INK,
    )

    # Web card
    add_round(s, Emu(502920), Emu(1680000), Emu(6400800), Emu(3650000), WHITE)
    tb(s, Emu(731520), Emu(1780000), Emu(5943600), Emu(380000), "Веб-версия", size=20, bold=True, color=TEAL)
    tb(
        s,
        Emu(731520),
        Emu(2140000),
        Emu(5943600),
        Emu(420000),
        "С готовым контентом. Цена за 1 работника.",
        size=14,
        color=MUTED,
    )
    add_price_table(s, Emu(731520), Emu(2580000), Emu(5943600), Emu(2520000), WEB_PRICES)

    # Server card
    add_round(s, Emu(7100640), Emu(1680000), Emu(4557960), Emu(3650000), WHITE)
    tb(s, Emu(7328040), Emu(1780000), Emu(4102560), Emu(380000), "Серверная версия", size=20, bold=True, color=NAVY)
    tb(
        s,
        Emu(7328040),
        Emu(2160000),
        Emu(4102560),
        Emu(420000),
        "Размещение на инфраструктуре заказчика",
        size=14,
        color=MUTED,
    )
    bullets(
        s,
        Emu(7328040),
        Emu(2620000),
        Emu(4102560),
        Emu(2400000),
        [
            "Стоимость в этом предложении не оговаривается",
            "Условия поставки и внедрения — отдельно",
            "Контур, доступы и сопровождение — под вашу инфраструктуру",
        ],
        size=15,
        color=INK,
        spacing=14,
    )

    add_round(s, Emu(502920), Emu(5450000), Emu(11155680), Emu(980000), WHITE)
    tb(
        s,
        Emu(731520),
        Emu(5580000),
        Emu(10700000),
        Emu(720000),
        "Дополнительно могут быть включены услуги по разработке или адаптации контента под специфику компании. Стоимость определяется по техническому заданию.",
        size=15,
        color=INK,
    )

    tb(s, Emu(10424160), Emu(6565392), Emu(1280160), Emu(256032), "17  /  17", size=11, color=FOOT, align=PP_ALIGN.RIGHT)

    try:
        ns = s.notes_slide
        tf = ns.notes_text_frame
        if tf is not None:
            tf.text = (
                "Цены — из коммерческого предложения СМАРТА на web-версию с готовым контентом. "
                "Серверную стоимость не называть. Контент под компанию — отдельной строкой по ТЗ."
            )
    except Exception:
        pass
    return s


def main():
    prs = Presentation(SRC)
    # Idempotent: if a commerce last slide is already there, skip adding another.
    last = prs.slides[-1]
    last_text = " ".join(sh.text_frame.text for sh in last.shapes if sh.has_text_frame)
    if "Веб-версия и серверная версия" in last_text:
        print("commerce slide already present, not duplicating")
    else:
        build_slide(prs)
    retitle_footers(prs, len(prs.slides))
    prs.save(SRC)
    print("saved", SRC, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
