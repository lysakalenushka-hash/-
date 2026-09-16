#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Собрать zip LMS 213–219: теплопотребители по системам (вентиляция / отопление / технооборудование)."""

from __future__ import annotations

import shutil
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RTN_PATH = Path("Приложение_2_РТН.xlsx")
KEYS_PATH = Path("Вентиляция_и_отопление_тест.xlsx")
TEMPLATE = Path("тесты_электробезопасность_РТН/199_neprom_dovyshe1000_G5.xlsx")
OUTDIR = Path("тесты_LMS_213-219")
ZIP_NAME = Path("teplopotrebiteli_lms_213-219.zip")

FONT = Font(name="Times New Roman", size=11)
ALIGN = Alignment(wrap_text=True, vertical="top")
HEAD = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="1F4E79")

PROGRAMS = [
    (213, {"v"}, "Теплопотребители. Вентиляция"),
    (214, {"h"}, "Теплопотребители. Отопление"),
    (215, {"t"}, "Теплопотребители. Технооборудование"),
    (216, {"v", "h"}, "Теплопотребители. Вентиляция и отопление"),
    (217, {"t", "v"}, "Теплопотребители. Технооборудование и вентиляция"),
    (218, {"t", "h"}, "Теплопотребители. Технооборудование и отопление"),
    (219, {"t", "h", "v"}, "Теплопотребители. Технооборудование, отопление и вентиляция"),
]


def parse_rtn_iii() -> list[dict]:
    wb = load_workbook(RTN_PATH, data_only=True)
    ws = wb["III"]
    qs: list[dict] = []
    cur = None
    npa = ""
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a == "№":
            continue
        is_num = isinstance(a, int) or (isinstance(a, str) and str(a).strip().isdigit())
        if a and not is_num:
            npa = str(a).strip()
            continue
        if is_num and b:
            if cur:
                qs.append(cur)
            cur = {
                "seq": len(qs) + 1,
                "n": int(str(a).strip()),
                "npa": npa,
                "q": str(b).strip(),
                "ans": [],
            }
            continue
        if cur and b and (a is None or a == ""):
            cur["ans"].append(str(b).strip())
    if cur:
        qs.append(cur)
    return qs


def load_keys(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Questions"]
    by: dict[str, dict] = {}
    cur = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        typ, txt, flag = row[2], row[3], row[4]
        if typ == "Вопрос":
            if cur:
                by[cur["q"]] = cur
            cur = {"q": str(txt or "").strip(), "ans": [], "flags": [], "comment": row[6]}
        elif typ == "Ответ" and cur is not None:
            cur["ans"].append(str(txt or "").strip())
            try:
                fv = int(float(flag)) if flag not in (None, "") else 0
            except (TypeError, ValueError):
                fv = 0
            cur["flags"].append(1 if fv == 1 else 0)
    if cur:
        by[cur["q"]] = cur
    return by


def pools(q: dict) -> set[str]:
    """Вентиляция / отопление / технооборудование по блокам ПТЭ 511 внутри листа III."""
    npa = q["npa"]
    if "охране труда" in npa or "924н" in npa or "первой помощи" in npa or "220н" in npa:
        return {"v", "h", "t"}
    n = q["n"]
    if 1 <= n <= 129 or 237 <= n <= 243:
        return {"v", "h", "t"}
    if 130 <= n <= 187:
        return {"t"}
    if 188 <= n <= 209:
        return {"h"}
    if 210 <= n <= 236:
        return {"v", "h"}
    return {"v", "h", "t"}


def file_name(lms: int, need: set[str]) -> str:
    parts = []
    if "v" in need:
        parts.append("vent")
    if "h" in need:
        parts.append("heat")
    if "t" in need:
        parts.append("tech")
    return f"{lms}_{'_'.join(parts)}.xlsx"


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def write_lms(template: Path, dest: Path, program: str, items: list[dict]) -> None:
    shutil.copy2(template, dest)
    wb = load_workbook(dest)
    ws = wb["Questions"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for num, it in enumerate(items, 1):
        flags = it["flags"]
        n_ok = sum(flags) or 1
        qvals = [program, num, "Вопрос", it["q"], n_ok, "Выбор ответа", it.get("comment"), None, program]
        for c, v in enumerate(qvals, 1):
            cell = ws.cell(row, c, v)
            cell.font = FONT
            cell.alignment = ALIGN
        ws.row_dimensions[row].height = 36
        row += 1
        for text, flag in zip(it["ans"], flags):
            avals = [program, num, "Ответ", text, int(flag)]
            for c, v in enumerate(avals, 1):
                cell = ws.cell(row, c, v)
                cell.font = FONT
                cell.alignment = ALIGN
            ws.row_dimensions[row].height = 22
            row += 1
    ws_p = wb["Теги программ"]
    ws_p.cell(2, 1).value = program
    ws_p.cell(2, 2).value = program
    wb.save(dest)


def write_table(path: Path, packs: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Вопросы"
    headers = ["LMS", "Программа", "№", "НПА", "Вопрос", "Системы", "Правильные ответы", "Ключ 1/0"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HEAD
        cell.fill = FILL_H
        cell.alignment = ALIGN
    row = 2
    for pack in packs:
        for i, it in enumerate(pack["items"], 1):
            correct = [a for a, f in zip(it["ans"], it["flags"]) if f]
            ws.cell(row, 1, pack["lms"]).font = FONT
            ws.cell(row, 2, pack["program"]).font = FONT
            ws.cell(row, 2).alignment = ALIGN
            ws.cell(row, 3, i).font = FONT
            ws.cell(row, 4, it["npa"]).font = FONT
            ws.cell(row, 4).alignment = ALIGN
            ws.cell(row, 5, it["q"]).font = FONT
            ws.cell(row, 5).alignment = ALIGN
            ws.cell(row, 6, it["systems"]).font = FONT
            ws.cell(row, 7, " | ".join(correct)).font = FONT
            ws.cell(row, 7).alignment = ALIGN
            ws.cell(row, 8, ",".join(str(int(f)) for f in it["flags"])).font = FONT
            ws.row_dimensions[row].height = 40
            row += 1
    widths = [8, 62, 8, 40, 70, 18, 70, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Сводка")
    ws2.append(["LMS", "Программа", "Вопросов", "С ключом", "В LMS (скриншот)"])
    ui = {213: 87, 214: 87, 215: 92, 216: 93, 217: 100, 218: 100, 219: 106}
    for pack in packs:
        k = sum(1 for it in pack["items"] if sum(it["flags"]) > 0)
        ws2.append([pack["lms"], pack["program"], len(pack["items"]), k, ui[pack["lms"]]])
    ws2.append([])
    ws2.append(["Как разделены блоки листа III РТН"])
    ws2.append(["Общие (все программы 213–219)", "ПОТ 924н, первая помощь 220н, ПТЭ 511 п.1–129 и теплопотребляющие установки п.237–243"])
    ws2.append(["Вентиляция + отопление", "Тепловые пункты и теплообменники ПТЭ п.210–236"])
    ws2.append(["Только отопление (дополнительно)", "Тепловые сети и насосные ПТЭ п.188–209"])
    ws2.append(["Только технооборудование (дополнительно)", "Топливо, котлы, сосуды, ВХР, баки ПТЭ п.130–187"])
    wb.save(path)


def attach_keys(qs: list[dict], bank: dict[str, dict]) -> list[dict]:
    items = []
    for q in qs:
        src = bank.get(q["q"])
        flags = list(src["flags"]) if src and len(src.get("flags") or []) == len(q["ans"]) else [0] * len(q["ans"])
        comment = q["npa"]
        if src and src.get("comment"):
            comment = str(src["comment"])
        sys = pools(q)
        names = []
        if "v" in sys:
            names.append("В")
        if "h" in sys:
            names.append("О")
        if "t" in sys:
            names.append("Т")
        items.append(
            {
                "q": q["q"],
                "ans": q["ans"],
                "flags": flags,
                "comment": comment,
                "npa": q["npa"],
                "systems": "".join(names),
                "sys": sys,
            }
        )
    return items


def main() -> None:
    rtn = parse_rtn_iii()
    bank = load_keys(KEYS_PATH)
    all_items = attach_keys(rtn, bank)

    if OUTDIR.exists():
        shutil.rmtree(OUTDIR)
    OUTDIR.mkdir()

    packs = []
    for lms, need, program in PROGRAMS:
        items = [it for it in all_items if it["sys"] & need]
        dest = OUTDIR / file_name(lms, need)
        write_lms(TEMPLATE, dest, program, items)
        keyed = sum(1 for it in items if sum(it["flags"]) > 0)
        packs.append({"lms": lms, "program": program, "items": items, "file": dest.name})
        print(f"{lms} {dest.name} Q={len(items)} keyed={keyed}")

    write_lms(TEMPLATE, OUTDIR / "000_все_категории_LMS.xlsx", "ВСЕ КАТЕГОРИИ 213-219", [])
    shutil.copy2(TEMPLATE, OUTDIR / "000_все_категории_LMS.xlsx")
    wb = load_workbook(OUTDIR / "000_все_категории_LMS.xlsx")
    ws = wb["Questions"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    prev = None
    num = 0
    combined_items = []
    for pack in packs:
        for it in pack["items"]:
            combined_items.append({**it, "program": pack["program"]})
    for it in combined_items:
        if it["program"] != prev:
            prev = it["program"]
            num = 0
        num += 1
        flags = it["flags"]
        n_ok = sum(flags) or 1
        qvals = [it["program"], num, "Вопрос", it["q"], n_ok, "Выбор ответа", it.get("comment"), None, it["program"]]
        for c, v in enumerate(qvals, 1):
            cell = ws.cell(row, c, v)
            cell.font = FONT
            cell.alignment = ALIGN
        ws.row_dimensions[row].height = 36
        row += 1
        for text, flag in zip(it["ans"], flags):
            avals = [it["program"], num, "Ответ", text, int(flag)]
            for c, v in enumerate(avals, 1):
                cell = ws.cell(row, c, v)
                cell.font = FONT
                cell.alignment = ALIGN
            ws.row_dimensions[row].height = 22
            row += 1
    ws_p = wb["Теги программ"]
    if ws_p.max_row > 1:
        ws_p.delete_rows(2, ws_p.max_row)
    for i, pack in enumerate(packs, 2):
        ws_p.cell(i, 1).value = pack["program"]
        ws_p.cell(i, 2).value = pack["program"]
    wb.save(OUTDIR / "000_все_категории_LMS.xlsx")

    write_table(OUTDIR / "000_все_категории_таблица.xlsx", packs)
    (OUTDIR / "README.txt").write_text(
        "Тесты теплопотребителей LMS 213–219\n"
        "Вопросы — приложение 2 РТН, лист III (276), ключи из Вентиляция_и_отопление_тест.xlsx.\n"
        "Разделение по системам: общие блоки во все программы; тепловые сети — отопление;\n"
        "топливо/котлы/сосуды — технооборудование; тепловые пункты — вентиляция и отопление.\n"
        "В официальном листе III нет колонок «+» как у электробезопасности, поэтому это разбиение по главам ПТЭ 511.\n"
        "В LMS на скриншоте вопросов меньше (87–106) — в файлах полный лист III с фильтром систем.\n",
        encoding="utf-8",
    )

    if ZIP_NAME.exists():
        ZIP_NAME.unlink()
    with zipfile.ZipFile(ZIP_NAME, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(OUTDIR.iterdir()):
            zf.write(p, f"{OUTDIR.name}/{p.name}")
    print("zip", ZIP_NAME, ZIP_NAME.stat().st_size)


if __name__ == "__main__":
    main()
