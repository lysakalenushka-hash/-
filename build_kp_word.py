#!/usr/bin/env python3
"""Сборка коммерческого предложения SMARTA в формате Word."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "КП_Обучение Тюмень 22.10.2026.docx"

FONT = "Times New Roman"
BLUE = RGBColor(0x4F, 0xAB, 0xD9)


def set_run_font(run, size=12, italic=False, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)


def set_paragraph_format(
    p,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before=0,
    space_after=0,
    line=14.9,
    first_line=None,
    left=0,
    right=0,
    hanging=None,
):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    pf.left_indent = Cm(left) if not isinstance(left, (Pt, Twips, Emu, Cm)) else left
    pf.right_indent = Cm(right) if not isinstance(right, (Pt, Twips, Emu, Cm)) else right
    if hanging is not None:
        pf.first_line_indent = -Cm(hanging) if not isinstance(hanging, (Pt, Twips, Emu, Cm)) else -hanging
    elif first_line is not None:
        pf.first_line_indent = (
            Cm(first_line) if not isinstance(first_line, (Pt, Twips, Emu, Cm)) else first_line
        )
    else:
        pf.first_line_indent = Cm(0)


def add_text(p, text, size=12, italic=False, bold=False):
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, bold=bold)
    return run


def disable_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for key, val in kwargs.items():
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_fixed(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w.twips if hasattr(w, "twips") else w)))
        grid.append(gc)
    tbl.insert(1, grid)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def add_float_picture(paragraph, image_path, width, left, top):
    """Вставляет подпись как плавающую картинку поверх текста."""
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=width)._inline
    drawing = inline.getparent()
    extent = inline.find(qn("wp:extent"))
    cx, cy = extent.get("cx"), extent.get("cy")
    docPr = inline.find(qn("wp:docPr"))
    docPr_id = docPr.get("id")
    docPr_name = docPr.get("name")

    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "114300")
    anchor.set("distR", "114300")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "251658240")
    anchor.set("behindDoc", "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "column")
    pos_h_off = OxmlElement("wp:posOffset")
    pos_h_off.text = str(int(left))
    pos_h.append(pos_h_off)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "paragraph")
    pos_v_off = OxmlElement("wp:posOffset")
    pos_v_off.text = str(int(top))
    pos_v.append(pos_v_off)
    anchor.append(pos_v)

    ext = OxmlElement("wp:extent")
    ext.set("cx", cx)
    ext.set("cy", cy)
    anchor.append(ext)

    effect = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        effect.set(k, "0")
    anchor.append(effect)

    wrap = OxmlElement("wp:wrapNone")
    anchor.append(wrap)

    docpr = OxmlElement("wp:docPr")
    docpr.set("id", docPr_id)
    docpr.set("name", docPr_name)
    anchor.append(docpr)

    graphic = inline.find(qn("a:graphic"))
    drawing.remove(inline)
    drawing.append(anchor)
    anchor.append(graphic)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # Как в исходном PDF: левое ~3.0 см, правое ~1.4 см
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.4)
    section.top_margin = Cm(0.15)
    section.bottom_margin = Cm(1.4)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0.5)

    # Шапка: вылет к левому краю страницы, как в исходном PDF
    header_p = doc.add_paragraph()
    pf = header_p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.left_indent = Cm(-3.0)
    pf.first_line_indent = Cm(0)
    run = header_p.add_run()
    run.add_picture(str(ASSETS / "header.png"), width=Cm(10.85))

    # Адресат справа
    addressee = [
        "Начальнику отдела",
        "по развитию компетенций",
        "в области производственной безопасности",
        "Управления производственной безопасности",
        "",
        "ПАО «Газпром нефть»",
        "",
    ]
    for i, line in enumerate(addressee):
        p = doc.add_paragraph()
        set_paragraph_format(
            p,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            line=13.8,
            space_before=0,
            space_after=0,
        )
        if line:
            add_text(p, line)

    # № слева и ФИО справа в одной строке
    row = doc.add_paragraph()
    set_paragraph_format(row, align=WD_ALIGN_PARAGRAPH.LEFT, line=13.8, space_after=16)
    tab_stops = row.paragraph_format.tab_stops
    usable = section.page_width - section.left_margin - section.right_margin
    tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    add_text(row, "№ 217/26 от 25.09.2026")
    add_text(row, "\tГаллямовой А.К.")

    # Обращение
    greet = doc.add_paragraph()
    set_paragraph_format(
        greet,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line=14.9,
        space_before=2,
        space_after=8,
    )
    add_text(greet, "Уважаемая Айгуль Камиловна!")

    body_paras = [
        "22 октября 2026 года в г. Тюмень состоится корпоративное мероприятие ПАО «Газпром нефть», посвященное вопросам обучения работников по охране труда.",
        "Компания «СМАРТА» предлагает Вам рассмотреть возможность участия в указанном мероприятии Титовой Алены Петровны, члена рабочей группы Минтруда России по разработке Правил обучения по охране труда и проверки знания требований охраны труда (постановление Правительства Российской Федерации от 24 декабря 2021 г. № 2464), руководителя проекта по разработке корпоративного стандарта ПАО «Газпром нефть», устанавливающего требования к процедуре обучения работников по охране труда, включая оказание следующих услуг:",
    ]
    for text in body_paras:
        p = doc.add_paragraph()
        set_paragraph_format(
            p,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line=1.25,
            line=14.9,
            space_after=8,
        )
        add_text(p, text)

    # Пункт 1
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left=1.88,
        hanging=0.63,
        line=14.9,
        space_after=4,
    )
    add_text(
        p,
        "1) модерация корпоративного мероприятия по вопросам обучения (управление дискуссией, экспертная поддержка, подведение итогов);",
    )

    # Пункт 2
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left=1.88,
        hanging=0.63,
        line=14.9,
        space_after=2,
    )
    add_text(
        p,
        "2) проведение обучающей сессии с разбором требований корпоративного стандарта:",
    )

    bullets = [
        "логика построения процедур инструктажа, обучения и стажировки;",
        "подходы, использованные при разработке отдельных положений;",
        "разбор замечаний к тексту стандарта, поступивших в процессе согласования, с обоснованием принятых решений;",
        "практические аспекты применения стандарта в деятельности дочерних обществ;",
        "ответы на вопросы участников.",
    ]
    for item in bullets:
        p = doc.add_paragraph()
        set_paragraph_format(
            p,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            left=2.52,
            hanging=0.63,
            line=15.7,
            space_after=0,
        )
        add_text(p, "•  " + item)

    after = [
        "По итогам обучающей сессии участники получат единое понимание требований корпоративного стандарта, принципов реализации основных процедур обучения и их внедрения в деятельность организаций Группы компаний ПАО «Газпром нефть».",
        "Формат мероприятия предполагает сочетание модерации, экспертного обучения и предметного обсуждения, что позволит обеспечить содержательное взаимодействие участников и подготовить основу для последующего практического применения корпоративного стандарта.",
        "Стоимость услуг составляет 125 000,00 руб., включая командировочные и иные расходы, а также НДС 5% – 5 952,38 руб.",
    ]
    for i, text in enumerate(after):
        p = doc.add_paragraph()
        set_paragraph_format(
            p,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line=1.25,
            line=14.9,
            space_before=8 if i == 0 else 6,
            space_after=2,
        )
        add_text(p, text)

    # Подпись
    respect = doc.add_paragraph()
    set_paragraph_format(
        respect,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        line=14.9,
        space_before=18,
        space_after=8,
    )
    add_text(respect, "С уважением,")

    sig = doc.add_paragraph()
    set_paragraph_format(sig, align=WD_ALIGN_PARAGRAPH.LEFT, line=14.9, space_after=0)
    tab_stops = sig.paragraph_format.tab_stops
    tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    add_text(sig, "генеральный директор")
    add_text(sig, "\tА.В. Москвичев")
    add_float_picture(
        sig,
        ASSETS / "signature.png",
        width=Cm(3.5),
        left=Emu(Cm(6.6)),
        top=Emu(Cm(-1.25)),
    )

    # Свойства файла
    core = doc.core_properties
    core.author = "ООО «СМАРТА»"
    core.title = "КП_Обучение Тюмень 22.10.2026"
    core.subject = "Коммерческое предложение № 217/26 от 25.09.2026"
    core.comments = "Участие Титовой А.П. в корпоративном мероприятии ПАО «Газпром нефть» 22.10.2026, г. Тюмень"

    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
