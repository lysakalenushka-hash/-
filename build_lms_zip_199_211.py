#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Собрать zip: 13 тестов LMS 199–211 + общий файл."""

from __future__ import annotations

import shutil
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_elektro_v_trainer_keys import (
    FONT as KEY_FONT,
    best_match,
    load_trainers,
    map_correct,
    parse_rtn_v,
)

RTN_PATH = Path("Приложение_2_РТН.xlsx")
TEMPLATE = Path("тесты_электробезопасность_РТН/199_neprom_dovyshe1000_G5.xlsx")
OUTDIR = Path("тесты_LMS_199-211")
ZIP_NAME = Path("elektrobez_lms_199-211.zip")

FONT = Font(name="Times New Roman", size=11)
ALIGN = Alignment(wrap_text=True, vertical="top")
HEAD = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="1F4E79")

PROGRAMS = [
    # lms, kind, above, group, screenshot name
    (199, "непромышленные", True, 5, "Потребители электроэнергии. Непромышленные. До и выше 1000 В V Группа"),
    (200, "непромышленные", False, 2, "Потребители электроэнергии. Непромышленные. До 1000 В II Группа"),
    (201, "непромышленные", False, 3, "Потребители электроэнергии. Непромышленные. До 1000 В III Группа"),
    (202, "непромышленные", False, 4, "Потребители электроэнергии. Непромышленные. До 1000 В IV Группа"),
    (203, "непромышленные", False, 5, "Потребители электроэнергии. Непромышленные. До 1000 В V Группа"),
    (204, "промышленные", True, 2, "Потребители электроэнергии. Промышленные. До и выше 1000 В II Группа"),
    (205, "промышленные", True, 3, "Потребители электроэнергии. Промышленные. До и выше 1000 В III Группа"),
    (206, "промышленные", True, 4, "Потребители электроэнергии. Промышленные. До и выше 1000 В IV Группа"),
    (207, "промышленные", True, 5, "Потребители электроэнергии. Промышленные. До и выше 1000 В V Группа"),
    (208, "промышленные", False, 2, "Потребители электроэнергии. Промышленные. До 1000 В II Группа"),
    (209, "промышленные", False, 3, "Потребители электроэнергии. Промышленные. До 1000 В III Группа"),
    (210, "промышленные", False, 4, "Потребители электроэнергии. Промышленные. До 1000 В IV Группа"),
    (211, "промышленные", False, 5, "Потребители электроэнергии. Промышленные. До 1000 В V Группа"),
]


def parse_sheet(ws):
    npa = None
    questions = []
    cur = None
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a == "№":
            continue
        is_num = isinstance(a, int) or (isinstance(a, str) and str(a).strip().isdigit())
        if a and not is_num:
            npa = str(a).strip()
            continue
        if is_num and b:
            if cur:
                questions.append(cur)

            def plus(c):
                return ws.cell(r, c).value == "+"

            cur = {
                "q": str(b).strip(),
                "ans": [],
                "npa": npa or "",
                "u1": plus(3),
                "u2": plus(4),
                "g2": plus(5),
                "g3": plus(6),
                "g4": plus(7),
                "g5": plus(8),
            }
            continue
        if cur and b and (a is None or a == ""):
            cur["ans"].append(str(b).strip())
    if cur:
        questions.append(cur)
    return questions


def filt(qs, above: bool, group: int):
    key = f"g{group}"
    out = []
    for q in qs:
        if not q[key]:
            continue
        if above and q["u2"]:
            out.append(q)
        if (not above) and q["u1"]:
            out.append(q)
    return out


def file_name(lms, kind, above, group):
    volt = "dovyshe1000" if above else "do1000"
    k = "prom" if kind == "промышленные" else "neprom"
    return f"{lms}_{k}_{volt}_G{group}.xlsx"


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def write_lms(template: Path, dest: Path, program: str, items: list[dict]):
    shutil.copy2(template, dest)
    wb = load_workbook(dest)
    ws = wb["Questions"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for num, it in enumerate(items, 1):
        flags = it["flags"]
        n_ok = sum(flags) or 1
        qvals = [program, num, "Вопрос", it["q"], n_ok, "Выбор ответа", it.get("comment"), None, program]
        for c, v in enumerate(qvals, 1):
            cell = ws.cell(row, c, v)
            cell.font = FONT
            cell.alignment = ALIGN
        ws.row_dimensions[row].height = 36
        row += 1
        for text, flag in zip(it["ans"], flags):
            avals = [program, num, "Ответ", text, int(flag)]
            for c, v in enumerate(avals, 1):
                cell = ws.cell(row, c, v)
                cell.font = FONT
                cell.alignment = ALIGN
            ws.row_dimensions[row].height = 22
            row += 1
    ws_p = wb["Теги программ"]
    ws_p.cell(2, 1).value = program
    ws_p.cell(2, 2).value = program
    wb.save(dest)


_CACHE = {}


def keyed_items(qs, bank):
    items = []
    keyed = 0
    for q in qs:
        ck = q["q"]
        if ck in _CACHE:
            flags, comment = _CACHE[ck]
        else:
            hit, sc, how = best_match(q["q"], bank)
            flags = map_correct(q["ans"], hit["correct"] if hit else [])
            if sum(flags) > 0:
                comment = f"{q['npa']}. Ключ: {hit.get('src', '')} ({how})"
            else:
                comment = f"{q['npa']}. Ключ не проставлен"
            _CACHE[ck] = (flags, comment)
        if sum(flags) > 0:
            keyed += 1
        items.append({"q": q["q"], "ans": q["ans"], "flags": flags, "comment": comment, "npa": q["npa"]})
    return items, keyed


def write_combined_human(path: Path, packs: list[dict]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Все категории"
    headers = ["LMS", "Программа", "№", "НПА", "Вопрос", "Правильные ответы", "Ключ 1/0 по вариантам"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HEAD
        cell.fill = FILL_H
        cell.alignment = ALIGN
    row = 2
    for pack in packs:
        for i, it in enumerate(pack["items"], 1):
            correct = [a for a, f in zip(it["ans"], it["flags"]) if f]
            ws.cell(row, 1, pack["lms"]).font = FONT
            ws.cell(row, 2, pack["program"]).font = FONT
            ws.cell(row, 2).alignment = ALIGN
            ws.cell(row, 3, i).font = FONT
            ws.cell(row, 4, it["npa"]).font = FONT
            ws.cell(row, 4).alignment = ALIGN
            ws.cell(row, 5, it["q"]).font = FONT
            ws.cell(row, 5).alignment = ALIGN
            ws.cell(row, 6, " | ".join(correct)).font = FONT
            ws.cell(row, 6).alignment = ALIGN
            ws.cell(row, 7, ",".join(str(int(f)) for f in it["flags"])).font = FONT
            ws.row_dimensions[row].height = 40
            row += 1
    widths = [8, 62, 8, 36, 70, 70, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Сводка")
    ws2.append(["LMS", "Программа", "Вопросов", "С ключом"])
    for pack in packs:
        k = sum(1 for it in pack["items"] if sum(it["flags"]) > 0)
        ws2.append([pack["lms"], pack["program"], len(pack["items"]), k])
    wb.save(path)


def main():
    rtn_wb = load_workbook(RTN_PATH, data_only=True)
    industrial = parse_sheet(rtn_wb["V"])
    nonind = parse_sheet(rtn_wb["Vн"])
    bank, _ = load_trainers()

    if OUTDIR.exists():
        shutil.rmtree(OUTDIR)
    OUTDIR.mkdir()

    packs = []
    all_items = []
    for lms, kind, above, group, program in PROGRAMS:
        src = nonind if kind == "непромышленные" else industrial
        qs = filt(src, above, group)
        items, keyed = keyed_items(qs, bank)
        dest = OUTDIR / file_name(lms, kind, above, group)
        write_lms(TEMPLATE, dest, program, items)
        packs.append({"lms": lms, "program": program, "items": items, "file": dest.name})
        all_items.extend({"program": program, **it} for it in items)
        print(f"{lms} {dest.name} Q={len(items)} keyed={keyed}")

    write_lms(TEMPLATE, OUTDIR / "000_все_категории_LMS.xlsx", "ВСЕ КАТЕГОРИИ 199-211", [])
    # rewrite combined LMS with per-row folder = program
    shutil.copy2(TEMPLATE, OUTDIR / "000_все_категории_LMS.xlsx")
    wb = load_workbook(OUTDIR / "000_все_категории_LMS.xlsx")
    ws = wb["Questions"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    num = 0
    prev = None
    for it in all_items:
        if it["program"] != prev:
            prev = it["program"]
            num = 0
        num += 1
        flags = it["flags"]
        n_ok = sum(flags) or 1
        qvals = [it["program"], num, "Вопрос", it["q"], n_ok, "Выбор ответа", it.get("comment"), None, it["program"]]
        for c, v in enumerate(qvals, 1):
            cell = ws.cell(row, c, v)
            cell.font = FONT
            cell.alignment = ALIGN
        ws.row_dimensions[row].height = 36
        row += 1
        for text, flag in zip(it["ans"], flags):
            avals = [it["program"], num, "Ответ", text, int(flag)]
            for c, v in enumerate(avals, 1):
                cell = ws.cell(row, c, v)
                cell.font = FONT
                cell.alignment = ALIGN
            ws.row_dimensions[row].height = 22
            row += 1
    ws_p = wb["Теги программ"]
    # list all programs as tags
    ws_p.delete_rows(2, ws_p.max_row)
    for i, pack in enumerate(packs, 2):
        ws_p.cell(i, 1, pack["program"])
        ws_p.cell(i, 2, pack["program"])
    wb.save(OUTDIR / "000_все_категории_LMS.xlsx")

    write_combined_human(OUTDIR / "000_все_категории_таблица.xlsx", packs)

    readme = OUTDIR / "README.txt"
    readme.write_text(
        "Тесты электробезопасности LMS 199–211\n"
        "Вопросы — приложение 2 РТН (листы V и Vн).\n"
        "Ключи — тренажёры / дампы / спецразделы, см. электробезопасность_V_ключи_тренажёр.xlsx.\n\n"
        "000_все_категории_LMS.xlsx — все 13 программ в одном файле импорта (папка = название программы).\n"
        "000_все_категории_таблица.xlsx — человекочитаемая сводка.\n"
        "199_*.xlsx … 211_*.xlsx — отдельные категории.\n",
        encoding="utf-8",
    )

    if ZIP_NAME.exists():
        ZIP_NAME.unlink()
    with zipfile.ZipFile(ZIP_NAME, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(OUTDIR.iterdir()):
            zf.write(p, arcname=f"{OUTDIR.name}/{p.name}")
    print("zip", ZIP_NAME, ZIP_NAME.stat().st_size)


if __name__ == "__main__":
    main()
